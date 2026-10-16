"""Carga inicial de los accesorios de cada unidad: lo que dice su última responsiva (16-oct-2026).

    uv run --no-project --python 3.12 --with google-api-python-client --with google-auth-oauthlib --with openpyxl \
        python tools/accesorios/carga_inicial.py <paso> …

    plan <libro> [--excel <xlsx>]   SOLO LECTURA. Lee el Excel de Jorge (data/Responsivas - serie y entrega.xlsx), el libro
                                    (VEHICULOS, catálogo y movimientos de accesorios) y, en Drive, la carpeta 5.-RESPONSIVA
                                    de cada NUCO. Deja en .cache/accesorios/<fecha>/ plan.json y revisar.xlsx
    aplicar <plan.json> <libro>     ESCRIBE: agrega a MOVIMIENTOS ACCESORIOS VEHICULOS los movimientos CARGA INICIAL del plan.
                                    Bitácora (los IDs) en .cache/accesorios/bitacoras/. Después, en el editor de Apps Script
                                    de ESE proyecto: rehacerEstadoAccesorios() (la hoja ESTADO ACCESORIOS VEHICULOS)
    verificar <plan.json> <libro>   SOLO LECTURA: cada movimiento del plan está en la hoja, tal cual, una sola vez
    deshacer <bitácora> <libro>     borra de la hoja los renglones de esa bitácora (por ID; solo esos)

<libro>: 'dev' (el libro de pruebas del DEV de ayrton), 'prod' o un id. 'prod' solo se LEE: escribir ahí está en la lista
negra de tools/migracion/conexion.py y se quita a propósito, en su propio commit, el día que se pida.

Reglas (las mismas que el servicio, AccesoriosVehiculoService.gs):
  - La unidad es la del NUCO de la carpeta donde está la responsiva; el No. de serie del papel solo confirma (clave serie
    o NUCO, nunca FOLIO). Serie distinta → no se carga.
  - SI y NO son dato; (vacío), N/A y en blanco no generan movimiento (sin dato).
  - No se carga: formatos sin casillas, "no es responsiva", notas de Jorge sobre las marcas (rayada, fuera de la casilla,
    guiones), una responsiva que ya no es la vigente (está en RESPONSIVAS ANTERIORES: hay una más nueva) y una unidad que
    ya tiene movimientos (así el plan se puede repetir).
  - FECHA: la del texto del PDF (Accesorios leidos.xlsx); si no, cuándo se subió el PDF a Drive (createdTime, si es de antes de la
    homologación de NUCOS del 8-oct-2026); si tampoco, la de la carga, marcada en el plan. ORIGEN: la liga al PDF.
"""
import json

import random
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
sys.path.insert(0, str(AQUI.parent / "migracion"))
sys.path.insert(0, str(AQUI.parent / "nucos"))

CACHE = AQUI / ".cache"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
EXCEL = RAIZ / "data" / "Responsivas - serie y entrega.xlsx"
LEIDOS = RAIZ.parent / "Responsivas por NUCO" / "Accesorios leidos.xlsx"
LIBROS = {"dev": "1fC77Uu1ePVUySNvhgWXMHqWpLhGhBMTZZMEblU2nUhI", "prod": "17YrtuMYjGGg5LZm1BxMNrEZyCxTMI2DiUHSVUIldUlc"}
PROHIBIDOS = {LIBROS["prod"]: "PRODUCCIÓN (ControlVehicular)"}
HOJA_VEH, HOJA_CAT, HOJA_MOV = "VEHICULOS", "CATALOGO ACCESORIOS VEHICULOS", "MOVIMIENTOS ACCESORIOS VEHICULOS"
COLUMNAS_MOV = ["ID", "FECHA", "ID VEHICULO", "NUCCO", "ACCESORIO", "MOVIMIENTO", "TIENE", "MOTIVO", "ORIGEN", "EVIDENCIA", "REGISTRO"]
FORMATOS = ("Completo (13 casillas)", "Corto (5 casillas)")
DUDA_MARCAS = re.compile(r"MARCA|RAYAD|GUION|CASILLA (SI|NO)", re.I)
REGISTRO = "Carga inicial (responsivas leídas por Jorge)"
CARPETA = "application/vnd.google-apps.folder"
HOMOLOGACION = datetime(2026, 10, 8, tzinfo=timezone.utc)


def normal(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def id_libro(texto):
    return LIBROS.get(texto, texto)


def sheets():
    from autorizar import credenciales
    from googleapiclient.discovery import build
    return build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()


def leer_hoja(api, libro, hoja):
    """[{encabezado: valor}] (como texto, tal cual lo muestra la hoja)"""
    v = api.values().get(spreadsheetId=libro, range="'%s'!A1:ZZ" % hoja, valueRenderOption="FORMATTED_VALUE").execute().get("values", [])
    if not v:
        return [], []
    enc = [str(x).strip() for x in v[0]]
    return enc, [dict(zip(enc, r + [""] * (len(enc) - len(r)))) for r in v[1:]]


# ---- IDs como Ids.nuevo de src/utils/Ids.gs: PRE-<8 de tiempo base32><6 al azar> ----
ALFABETO = "0123456789ABCDEFGHJKMNPQRSTVWXYZ"
EPOCA_MS = 1577836800000


def base32(n, largo):
    s = ""
    while n > 0:
        s, n = ALFABETO[n % 32] + s, n // 32
    return s.rjust(largo, "0")


def ids_nuevos(prefijo, cuantos):
    azar = random.SystemRandom()
    t = int(time.time() * 1000) - EPOCA_MS
    # Un milisegundo distinto para cada uno: el orden del ID sigue el orden del plan (el desempate de estadoDe)
    return [prefijo + "-" + base32(t + i, 8) + "".join(ALFABETO[int(azar.random() * 32)] for _ in range(6)) for i in range(cuantos)]


def serie_cuadra(papel, hoja):
    """¿La serie leída del papel es la de la unidad? Igual, o los últimos 8 iguales (los '?' cuentan como cualquiera)."""
    a, b = re.sub(r"[^A-Z0-9?]", "", normal(papel)), re.sub(r"[^A-Z0-9]", "", normal(hoja))
    if not a:
        return None          # el papel no la trae: no confirma ni contradice
    if a == b:
        return True
    cola_a, cola_b = a[-8:], b[-8:]
    return len(cola_a) == 8 and len(cola_b) == 8 and all(x == "?" or x == y for x, y in zip(cola_a, cola_b))


def leer_excel(ruta):
    from openpyxl import load_workbook
    ws = load_workbook(ruta, read_only=True, data_only=True)["Responsivas"]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(x or "").strip() for x in filas[0]]
    return enc, [dict(zip(enc, r)) for r in filas[1:] if any(v not in (None, "") for v in r)]


def fechas_leidas():
    """ARCHIVO → fecha del texto del PDF (Accesorios leidos.xlsx), si está"""
    if not LEIDOS.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(LEIDOS, read_only=True, data_only=True)["Lectura"]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(x or "").strip() for x in filas[0]]
    i_a, i_f = enc.index("ARCHIVO"), next(i for i, c in enumerate(enc) if c.startswith("FECHA"))
    return {str(r[i_a]).strip(): r[i_f] for r in filas[1:] if r[i_a] and isinstance(r[i_f], datetime)}


def responsivas_en_drive(nucos):
    """NUCO → {'vigentes': {nombre: archivo}, 'anteriores': {nombre: archivo}} de su 1.-DOCUMENTACIÓN/5.-RESPONSIVA"""
    import expediente as e
    import reglas
    # Con createdTime: modifiedTime cambió al renombrar en la homologación de NUCOS (8 al 15-oct-2026)
    e.CAMPOS = e.CAMPOS + ",createdTime"
    d = e.drive()
    carpetas = {str(int(f["name"].strip())): f["id"] for f in e.hijos_de(d, [e.NUCOS])[e.NUCOS]
                if f["mimeType"] == CARPETA and f["name"].strip().isdigit()}
    ids = {n: carpetas[n] for n in nucos if n in carpetas}
    docs = {}
    for nid, hijos in e.hijos_de(d, list(ids.values())).items():
        for h in hijos:
            if h["mimeType"] == CARPETA and reglas.es_documentacion(h["name"]):
                docs[h["id"]] = nid
    resp = {}
    for did, hijos in e.hijos_de(d, list(docs)).items():
        for h in hijos:
            if h["mimeType"] == CARPETA and normal(h["name"]).replace(" ", "").startswith("5.-RESPONSIVA"):
                resp[h["id"]] = docs[did]
    nuco_de = {v: k for k, v in ids.items()}
    out = {n: {"vigentes": {}, "anteriores": {}} for n in nucos}
    anteriores = {}
    for rid, hijos in e.hijos_de(d, list(resp)).items():
        n = nuco_de[resp[rid]]
        for h in hijos:
            if h["mimeType"] == CARPETA:
                if normal(h["name"]) == reglas.RESPONSIVAS_ANTERIORES:
                    anteriores[h["id"]] = n
            else:
                out[n]["vigentes"][h["name"].strip()] = h
    for aid, hijos in e.hijos_de(d, list(anteriores)).items():
        for h in hijos:
            if h["mimeType"] != CARPETA:
                out[anteriores[aid]]["anteriores"][h["name"].strip()] = h
    return out


def plan(libro, ruta_excel):
    api = sheets()
    libro = id_libro(libro)
    _, vehiculos = leer_hoja(api, libro, HOJA_VEH)
    _, catalogo = leer_hoja(api, libro, HOJA_CAT)
    _, movimientos = leer_hoja(api, libro, HOJA_MOV)
    if not catalogo:
        raise SystemExit("El libro no tiene la hoja %s con piezas." % HOJA_CAT)
    con_movimientos = {m["ID VEHICULO"].strip() for m in movimientos if m.get("ID VEHICULO")}
    col_serie = next((c for c in (vehiculos[0] if vehiculos else {}) if normal(c) in ("SERIE VEHICULO", "NO SERIE", "SERIE")), None)
    por_nuco = {}
    for v in vehiculos:
        n = re.sub(r"\D", "", str(v.get("NUCCO", "")))
        if n:
            por_nuco.setdefault(str(int(n)), []).append(v)

    enc, filas = leer_excel(ruta_excel)
    # Cada columna de pieza del Excel → su CLAVE del catálogo (por el nombre de la pieza, sin acentos)
    por_nombre = {normal(p["NOMBRE"]): p["CLAVE"].strip() for p in catalogo}
    por_nombre.update({normal(p["CLAVE"]): p["CLAVE"].strip() for p in catalogo})
    piezas = {c: por_nombre[normal(c)] for c in enc if normal(c) in por_nombre}
    if len(piezas) != 13:
        raise SystemExit("Esperaba las 13 piezas de la responsiva en el Excel y encontré %d: %s" % (len(piezas), piezas))

    nucos = sorted({str(int(f["NUCO"])) for f in filas if str(f.get("NUCO") or "").strip().isdigit()}, key=int)
    print("Drive: 5.-RESPONSIVA de %d NUCO" % len(nucos), flush=True)
    drive = responsivas_en_drive(nucos)
    leidas = fechas_leidas()

    cargar, fuera = [], []
    for f in filas:
        nuco = str(int(f["NUCO"])) if str(f.get("NUCO") or "").strip().isdigit() else ""
        archivo = str(f.get("Archivo") or "").strip()
        nombre_pdf = re.sub(r"^\d+\s*-\s*", "", archivo)
        base = {"nuco": nuco, "archivo": archivo, "formato": f.get("Formato"), "nota": f.get("Nota / verificar") or "",
                "otros": f.get("Otros") or ""}
        valores = {piezas[c]: ("SI" if normal(f.get(c)) == "SI" else "NO" if normal(f.get(c)) == "NO" else "") for c in piezas}
        con_dato = {k: v for k, v in valores.items() if v}

        def no(motivo, **extra):
            fuera.append(dict(base, motivo=motivo, piezas=len(con_dato), **extra))

        if f.get("Formato") not in FORMATOS:
            no("formato sin casillas (%s)" % f.get("Formato"))
            continue
        if not con_dato:
            no("ninguna casilla marcada")
            continue
        if DUDA_MARCAS.search(str(base["nota"])):
            no("Jorge anotó duda en las marcas")
            continue
        unidades = por_nuco.get(nuco, [])
        if not unidades:
            no("el NUCO no está en VEHICULOS")
            continue
        if len(unidades) > 1:
            cuadran = [u for u in unidades if col_serie and serie_cuadra(f.get("No. de serie"), u.get(col_serie))]
            if len(cuadran) != 1:
                no("%d unidades con ese NUCO y la serie no decide" % len(unidades))
                continue
            unidades = cuadran
        u = unidades[0]
        cuadra = serie_cuadra(f.get("No. de serie"), u.get(col_serie)) if col_serie else None
        if cuadra is False:
            no("la serie del papel no es la de la unidad", serie_papel=str(f.get("No. de serie")), serie_hoja=u.get(col_serie, ""))
            continue
        id_veh = u["ID"].strip()
        if id_veh in con_movimientos:
            no("la unidad ya tiene movimientos (ya se cargó o se capturó)")
            continue
        en_drive = drive.get(nuco, {"vigentes": {}, "anteriores": {}})
        pdf = en_drive["vigentes"].get(nombre_pdf)
        if not pdf and nombre_pdf in en_drive["anteriores"]:
            no("ya no es la vigente (está en RESPONSIVAS ANTERIORES: hay una más nueva)")
            continue
        fecha, de_donde = leidas.get(archivo), "texto del PDF"
        if not fecha and pdf:
            creado = datetime.fromisoformat(pdf["createdTime"].replace("Z", "+00:00"))
            # Lo creado desde la homologación de NUCOS (copias de AppSheet, PDFs nuevos) no dice cuándo se firmó
            if creado < HOMOLOGACION:
                fecha, de_donde = creado, "cuando se subió a Drive"
        if not fecha:
            fecha, de_donde = datetime.now(timezone.utc), "SIN FECHA (la de la carga)"
        origen = ("https://drive.google.com/file/d/%s/view" % pdf["id"]) if pdf else archivo
        cargar.append(dict(base, id_vehiculo=id_veh, nucco=u.get("NUCCO", ""), folio=u.get("FOLIO", ""),
                           serie=("confirmada" if cuadra else "el papel no la trae"), fecha=fecha.strftime("%Y-%m-%d"),
                           fecha_de=de_donde, origen=origen, pdf_en_drive=bool(pdf), valores=con_dato))

    movs = [(c, clave, tiene) for c in cargar for clave, tiene in sorted(c["valores"].items(), key=lambda x: list(piezas.values()).index(x[0]))]
    ids = ids_nuevos("MAV", len(movs))
    filas_mov = [{"ID": i, "FECHA": c["fecha"], "ID VEHICULO": c["id_vehiculo"], "NUCCO": c["nucco"], "ACCESORIO": clave,
                  "MOVIMIENTO": "CARGA INICIAL", "TIENE": tiene, "MOTIVO": "Responsiva " + c["archivo"], "ORIGEN": c["origen"],
                  "EVIDENCIA": "", "REGISTRO": REGISTRO} for i, (c, clave, tiene) in zip(ids, movs)]

    carpeta = CACHE / time.strftime("%Y%m%d-%H%M%S")
    carpeta.mkdir(parents=True, exist_ok=True)
    datos = {"libro": libro, "excel": str(ruta_excel), "unidades": cargar, "fuera": fuera, "movimientos": filas_mov}
    (carpeta / "plan.json").write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    excel_revisar(carpeta / "revisar.xlsx", datos, list(piezas.values()))

    import collections
    print("Se cargan %d unidades → %d movimientos (SI %d · NO %d)" % (
        len(cargar), len(filas_mov), sum(m["TIENE"] == "SI" for m in filas_mov), sum(m["TIENE"] == "NO" for m in filas_mov)))
    print("  fecha:", dict(collections.Counter(c["fecha_de"] for c in cargar)))
    print("  serie:", dict(collections.Counter(c["serie"] for c in cargar)))
    print("No se cargan %d:" % len(fuera))
    for k, v in collections.Counter(re.sub(r" \(.*", "", x["motivo"]) for x in fuera).most_common():
        print("  %4d  %s" % (v, k))
    print("Plan:", carpeta / "plan.json")
    print("Para revisar:", carpeta / "revisar.xlsx")


def excel_revisar(ruta, datos, claves):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    a = wb.active
    a.title = "SE CARGAN"
    a.append(["NUCO", "FOLIO", "ID VEHICULO", "RESPONSIVA", "FECHA", "DE DÓNDE LA FECHA", "SERIE", "NOTA DE JORGE", "OTROS"] + claves)
    for c in sorted(datos["unidades"], key=lambda x: int(x["nuco"])):
        a.append([int(c["nuco"]), c["folio"], c["id_vehiculo"], c["archivo"], c["fecha"], c["fecha_de"], c["serie"], c["nota"],
                  c["otros"]] + [c["valores"].get(k, "") for k in claves])
    b = wb.create_sheet("NO SE CARGAN")
    b.append(["NUCO", "RESPONSIVA", "FORMATO", "MOTIVO", "PIEZAS CON DATO", "SERIE PAPEL", "SERIE HOJA", "NOTA DE JORGE", "OTROS"])
    for x in sorted(datos["fuera"], key=lambda x: (x["motivo"], int(x["nuco"] or 0))):
        b.append([int(x["nuco"]) if x["nuco"] else "", x["archivo"], x["formato"], x["motivo"], x["piezas"], x.get("serie_papel", ""),
                  x.get("serie_hoja", ""), x["nota"], x["otros"]])
    m = wb.create_sheet("MOVIMIENTOS")
    m.append(COLUMNAS_MOV)
    for r in datos["movimientos"]:
        m.append([r[c] for c in COLUMNAS_MOV])
    for ws in (a, b, m):
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


def exigir_escribible(libro):
    if libro in PROHIBIDOS:
        raise SystemExit("El libro %s es %s: esta herramienta no escribe ahí (ver PROHIBIDOS)." % (libro, PROHIBIDOS[libro]))


def aplicar(ruta_plan, libro):
    libro = id_libro(libro)
    exigir_escribible(libro)
    datos = json.loads(Path(ruta_plan).read_text(encoding="utf-8"))
    api = sheets()
    enc, movimientos = leer_hoja(api, libro, HOJA_MOV)
    if enc[:len(COLUMNAS_MOV)] != COLUMNAS_MOV:
        raise SystemExit("Las columnas de %s no son las esperadas: %s" % (HOJA_MOV, enc))
    # Bajo las mismas reglas del plan: si la unidad ya tiene movimientos (alguien capturó desde que se hizo el plan), no
    ya = {m["ID VEHICULO"].strip() for m in movimientos if m.get("ID VEHICULO")}
    ids_ya = {m["ID"].strip() for m in movimientos}
    nuevos = [r for r in datos["movimientos"] if r["ID VEHICULO"] not in ya and r["ID"] not in ids_ya]
    saltadas = {r["ID VEHICULO"] for r in datos["movimientos"]} - {r["ID VEHICULO"] for r in nuevos}
    if not nuevos:
        print("Nada que agregar (%d unidades ya tienen movimientos)." % len(saltadas))
        return
    bit = CACHE / "bitacoras" / (time.strftime("%Y%m%d-%H%M%S") + "-carga-inicial.jsonl")
    bit.parent.mkdir(parents=True, exist_ok=True)
    with open(bit, "w", encoding="utf-8") as b:
        b.write(json.dumps({"libro": libro, "plan": str(ruta_plan), "ids": [r["ID"] for r in nuevos]}, ensure_ascii=False) + "\n")
    # Escribe en bloques de 500 renglones, tal cual (la fecha yyyy-MM-dd la toma la hoja como fecha)
    for i in range(0, len(nuevos), 500):
        bloque = [[r[c] for c in COLUMNAS_MOV] for r in nuevos[i:i + 500]]
        api.values().append(spreadsheetId=libro, range="'%s'!A1" % HOJA_MOV, valueInputOption="USER_ENTERED",
                            insertDataOption="INSERT_ROWS", body={"values": bloque}).execute()
        print("  %d/%d" % (min(i + 500, len(nuevos)), len(nuevos)), flush=True)
    print("Agregados %d movimientos de %d unidades%s." % (len(nuevos), len({r["ID VEHICULO"] for r in nuevos}),
          (" (%d unidades saltadas: ya tenían movimientos)" % len(saltadas)) if saltadas else ""))
    print("Bitácora:", bit)
    print("Ahora, en el editor de Apps Script de ese proyecto: rehacerEstadoAccesorios() y revisarEstadoAccesorios().")


def verificar(ruta_plan, libro):
    libro = id_libro(libro)
    datos = json.loads(Path(ruta_plan).read_text(encoding="utf-8"))
    _, movimientos = leer_hoja(sheets(), libro, HOJA_MOV)
    por_id = {}
    for m in movimientos:
        por_id.setdefault(m["ID"].strip(), []).append(m)
    bien, mal = 0, []
    for r in datos["movimientos"]:
        h = por_id.get(r["ID"], [])
        if len(h) != 1:
            mal.append("%s (%s %s): %d renglones" % (r["ID"], r["ID VEHICULO"], r["ACCESORIO"], len(h)))
            continue
        h = h[0]
        distinto = [c for c in ("ID VEHICULO", "ACCESORIO", "MOVIMIENTO", "TIENE", "ORIGEN") if h.get(c, "").strip() != str(r[c])]
        if distinto:
            mal.append("%s: distinto en %s" % (r["ID"], ", ".join(distinto)))
        else:
            bien += 1
    print("bien: %d | mal: %d" % (bien, len(mal)))
    for x in mal[:30]:
        print("   ✘", x)
    return not mal


def deshacer(ruta_bitacora, libro):
    libro = id_libro(libro)
    exigir_escribible(libro)
    cab = json.loads(Path(ruta_bitacora).read_text(encoding="utf-8").splitlines()[0])
    if cab["libro"] != libro:
        raise SystemExit("La bitácora es del libro %s, no de %s." % (cab["libro"], libro))
    ids = set(cab["ids"])
    api = sheets()
    meta = api.get(spreadsheetId=libro, fields="sheets(properties(sheetId,title))").execute()
    sheet_id = next(s["properties"]["sheetId"] for s in meta["sheets"] if s["properties"]["title"] == HOJA_MOV)
    col = api.values().get(spreadsheetId=libro, range="'%s'!A:A" % HOJA_MOV).execute().get("values", [])
    filas = [i for i, v in enumerate(col) if i > 0 and v and v[0].strip() in ids]
    if not filas:
        print("Ninguno de los %d IDs está en la hoja." % len(ids))
        return
    # De abajo hacia arriba, para que borrar uno no recorra los que faltan
    pedidos = [{"deleteDimension": {"range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": i, "endIndex": i + 1}}}
               for i in sorted(filas, reverse=True)]
    for k in range(0, len(pedidos), 500):
        api.batchUpdate(spreadsheetId=libro, body={"requests": pedidos[k:k + 500]}).execute()
    print("Borrados %d de %d renglones de la bitácora." % (len(filas), len(ids)))
    print("Ahora: rehacerEstadoAccesorios() en el editor de ese proyecto.")


if __name__ == "__main__":
    a = sys.argv[1:]
    paso = a[0] if a else ""
    if paso == "plan" and len(a) >= 2:
        plan(a[1], Path(a[a.index("--excel") + 1]) if "--excel" in a else EXCEL)
    elif paso == "aplicar" and len(a) == 3:
        aplicar(a[1], a[2])
    elif paso == "verificar" and len(a) == 3:
        sys.exit(0 if verificar(a[1], a[2]) else 1)
    elif paso == "deshacer" and len(a) == 3:
        deshacer(a[1], a[2])
    else:
        sys.exit(__doc__)
