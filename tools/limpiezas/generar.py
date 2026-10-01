"""
Genera limpiezas-planeadas.xlsx (en la raíz del repo) a partir de datos.py.

    uv run --no-project --with openpyxl python tools/limpiezas/generar.py

El contenido se edita en datos.py, nunca en el Excel: lo que se cambie a mano en el Excel se
pierde la próxima vez que alguien lo genere. El formato es el mismo en todas las pestañas
(encabezado oscuro, texto ajustado arriba) y está aquí, no en los datos.
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from celda import Celda  # noqa: E402
from datos import HOJAS  # noqa: E402

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
DESTINO = os.path.join(RAIZ, 'limpiezas-planeadas.xlsx')

COLORES = {
    'amarillo': 'FFF2CC', 'azul': 'DDEBF7', 'verde': 'C6E0B4',
    'gris': 'E2E2E2', 'naranja': 'F8CBAD', 'rojo': 'FF9999',
}

# El color de cada estado de la pestaña Resumen. Un estado que no esté aquí truena, a propósito:
# así un "PENDENTE" mal escrito no pasa como un estado más.
ESTADOS = {
    'PENDIENTE': 'amarillo',
    'RE-MEDIR': 'azul',
    'EN PAUSA': 'azul',
    'HECHO': 'verde',
    'YA NO APLICA': 'gris',
    'DECISION DE NEGOCIO': 'naranja',
    'URGENTE': 'rojo',
}

TAMANO = 10
ENCABEZADO_FONDO = '1F4E5F'
ENCABEZADO_ALTO = 32


def relleno(nombre):
    return PatternFill(fill_type='solid', fgColor=COLORES[nombre])


def generar(destino=DESTINO):
    wb = Workbook()
    wb.remove(wb.active)
    for hoja in HOJAS:
        ws = wb.create_sheet(hoja['nombre'])
        enc = hoja['encabezados']
        for j, ancho in enumerate(hoja['anchos'], start=1):
            ws.column_dimensions[get_column_letter(j)].width = ancho

        for j, titulo in enumerate(enc, start=1):
            c = ws.cell(1, j, titulo)
            c.font = Font(bold=True, size=TAMANO, color='FFFFFF')
            c.fill = PatternFill(fill_type='solid', fgColor=ENCABEZADO_FONDO)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = ENCABEZADO_ALTO
        ws.freeze_panes = 'A2'

        col_estado = enc.index('Estado') if hoja['nombre'] == 'Resumen' else -1
        for i, fila in enumerate(hoja['filas'], start=2):
            if not fila:
                continue   # renglón vacío: separa secciones
            if len(fila) != len(enc):
                raise SystemExit('%s, fila %d: tiene %d celdas y la pestaña %d columnas'
                                 % (hoja['nombre'], i, len(fila), len(enc)))
            for j, x in enumerate(fila):
                celda = x if isinstance(x, Celda) else Celda(x)
                c = ws.cell(i, j + 1, celda.valor)
                c.font = Font(bold=celda.negrita, size=TAMANO)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                if j == col_estado:
                    if celda.valor not in ESTADOS:
                        raise SystemExit('%s, fila %d: estado desconocido %r (válidos: %s)'
                                         % (hoja['nombre'], i, celda.valor, ', '.join(ESTADOS)))
                    c.font = Font(bold=True, size=TAMANO)
                    c.fill = relleno(ESTADOS[celda.valor])
                elif celda.relleno:
                    c.fill = relleno(celda.relleno)
    wb.save(destino)
    return destino


if __name__ == '__main__':
    print('Generado: ' + generar(sys.argv[1] if len(sys.argv) > 1 else DESTINO))
