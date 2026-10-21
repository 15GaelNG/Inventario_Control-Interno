import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill

WIN_SCRATCH = sys.argv[1]
SALIDA = sys.argv[2]

CLAVES = ['GATO HIDRAULICO', 'GATO MECANICO', 'LLAVE CRUZ', 'MANERAL', 'LLANTA REFACCION',
          'KIT SEGURIDAD', 'PARASOL', 'TAPETES RUDO', 'TAPETES ALFOMBRA', 'CUBRE VOLANTE',
          'CUBRE ASIENTOS', 'BED LINER', 'CUBIERTA BATEA']

data = json.load(open(os.path.join(WIN_SCRATCH, 'compilado.json'), encoding='utf-8'))
filas = sorted(data['filas'], key=lambda f: (f['nuco'] is None, f['nuco']))
resumen = data['resumen']

wb = openpyxl.Workbook()

ws1 = wb.active
ws1.title = 'Resumen'
ws1['A1'] = 'Migración de accesorios desde Responsivas escaneadas'
ws1['A1'].font = Font(bold=True, size=14)
filas_resumen = [
    ('Total de archivos procesados', resumen['total_archivos']),
    ('Con datos de accesorios extraídos', resumen['con_datos']),
    ('Plantilla vieja (4 accesorios)', resumen['VIEJA_4']),
    ('Plantilla nueva (13 accesorios)', resumen['NUEVA_13']),
    ('No es una Responsiva (doc. distinto/ilegible)', resumen['NO_RESPONSIVA']),
    ('Excluidos (xlsx que no era responsiva)', resumen['EXCLUIDO_XLSX']),
    ('Ilegible / sin datos', resumen['sin_datos_ilegible']),
]
r = 3
for etiqueta, valor in filas_resumen:
    ws1.cell(row=r, column=1, value=etiqueta)
    ws1.cell(row=r, column=2, value=valor)
    r += 1
ws1.column_dimensions['A'].width = 48

ws2 = wb.create_sheet('Datos')
encabezados = ['NUCO', 'FOLIO', 'PLACA', 'Plantilla', 'Origen extracción', 'Archivo'] + CLAVES
ws2.append(encabezados)
for c in range(1, len(encabezados) + 1):
    ws2.cell(row=1, column=c).font = Font(bold=True)
ws2.freeze_panes = 'A2'

amarillo = PatternFill(start_color='FFFFF2CC', end_color='FFFFF2CC', fill_type='solid')
rojo = PatternFill(start_color='FFFCE4E4', end_color='FFFCE4E4', fill_type='solid')

for f in filas:
    if f['plantilla'] in ('NO_RESPONSIVA', 'XLSX_NO_RESPONSIVA') or f['datos'] is None:
        continue
    fila = [f['nuco'], f['folio'], f['placa'], f['plantilla'], f['origen_extraccion'], f['archivo']]
    for k in CLAVES:
        fila.append(f['datos'].get(k, ''))
    ws2.append(fila)

for col, width in zip('ABCDEF', [8, 12, 12, 12, 16, 40]):
    ws2.column_dimensions[col].width = width
for col in 'GHIJKLMNOPQRS':
    ws2.column_dimensions[col].width = 11

ws3 = wb.create_sheet('Revisar a mano')
ws3.append(['NUCO', 'FOLIO', 'Archivo', 'Motivo', 'Nota'])
for c in range(1, 6):
    ws3.cell(row=1, column=c).font = Font(bold=True)
for f in filas:
    if f['plantilla'] == 'NO_RESPONSIVA':
        motivo = 'El archivo no es una Responsiva Vehicular con checklist'
    elif f['plantilla'] == 'XLSX_NO_RESPONSIVA':
        motivo = 'Era un .xlsx que no es una responsiva (hoja de auditoría vieja)'
    elif f['datos'] is None:
        motivo = 'Ilegible / no se pudo leer con confianza'
    else:
        continue
    fila = ws3.append([f['nuco'], f['folio'], f['archivo'], motivo, f.get('nota') or ''])
for col, width in zip('ABCDE', [8, 12, 40, 45, 60]):
    ws3.column_dimensions[col].width = width

wb.save(SALIDA)
print('guardado:', SALIDA)
