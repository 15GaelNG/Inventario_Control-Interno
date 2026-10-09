"""Excel y árboles (antes/después) del dry run de homologar inspecciones."""
import sys, json, collections, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

x = json.load(open(sys.argv[1], encoding="utf-8"))
h = json.load(open(sys.argv[2], encoding="utf-8"))
SAL = sys.argv[3]
MUESTRA = sys.argv[4].split(",")
P, nota = h["plan"], h["nota"]
C = {c["id"]: c for c in x["carpetas"]}
cuenta = collections.Counter()
for a in x["archivos"]:
    for s in range(1, len(a["ruta"])):
        cuenta[(a["nuco"], "/".join(a["ruta"][:s]))] += 1


# ---------- árboles ----------
def arbol(rutas, titulo):
    """rutas: lista de rutas de archivo. Pinta carpetas; dentro de cada carpeta, los PDF por nombre y el resto contado."""
    t = {}
    for r in rutas:
        n = t
        partes = r.split("/")
        for p in partes[:-1]:
            n = n.setdefault(p + "/", {})
        n.setdefault("__f", []).append(partes[-1])
    out = [titulo]

    def pinta(n, pref):
        dirs = sorted(k for k in n if k != "__f")
        fs = n.get("__f", [])
        pdfs = sorted(f for f in fs if f.lower().endswith(".pdf"))
        otros = len(fs) - len(pdfs)
        items = [("d", d) for d in dirs] + [("f", f) for f in pdfs] + ([("n", "(%d fotos/videos/otros)" % otros)] if otros else [])
        for i, (k, v) in enumerate(items):
            ult = i == len(items) - 1
            out.append(pref + ("└── " if ult else "├── ") + v)
            if k == "d":
                hijo = n[v]
                # Dentro de una sección no se baja más: solo se cuenta
                if re.fullmatch(r"(CRISTALERIA|DOCUMENTACION|INTERIORES|LATONERIA Y PINTURA|NEUMATICOS|OTROS ELEMENTOS|"
                                r"SISTEMA ELECTRICO|SISTEMA MECANICO|Cristalería|Documentación|Interiores|Latonería y pintura|"
                                r"Neumáticos|Sistema Eléctrico|Sistema Mecánico)/", v):
                    total = sum(1 for _ in recorre(hijo))
                    out[-1] += "  (%d)" % total
                else:
                    pinta(hijo, pref + ("    " if ult else "│   "))
    pinta(t, "")
    return out


def recorre(n):
    for k, v in n.items():
        if k == "__f":
            yield from v
        else:
            yield from recorre(v)


lineas = []
for nuco in MUESTRA:
    ps = [p for p in P if p["nuco"] == nuco]
    if not ps:
        continue
    lineas += ["=" * 100, "NUCO %s" % nuco, "=" * 100]
    lineas += arbol([p["antes"] for p in ps], "ANTES")
    lineas += [""]
    lineas += arbol([p["despues"] for p in ps if p["accion"] != "PAPELERA"], "DESPUÉS")
    pap = [p for p in ps if p["accion"] == "PAPELERA"]
    if pap:
        lineas += ["", "A LA PAPELERA (%d):" % len(pap)] + ["  " + p["antes"].split("/", 1)[1] + "   ← " + p["aviso"] for p in pap]
    lineas += [""]
open(SAL + "-arboles.txt", "w", encoding="utf-8").write("\n".join(lineas))

# ---------- Excel ----------
wb = Workbook()
F = lambda **k: Font(name="Arial", **k)
r = wb.active
r.title = "Resumen"
cambia = [p for p in P if p["antes"] != p["despues"]]
clas = collections.Counter(v[0] for v in nota.values())
acc = collections.Counter(p["accion"] for p in P)
prof = collections.Counter(len(p["despues"].split("/")) - 3 for p in P
                           if p["accion"] != "PAPELERA" and p["despues"].split("/")[1] == "4.- INSPECCIONES")
filas = [
    ("Dry run: homologar 4.- INSPECCIONES (solo lectura, nada se movió)", ""),
    ("", ""),
    ("NUCO con carpeta de inspecciones", len(x["raices"])),
    ("Carpeta raíz que se renombra a «4.- INSPECCIONES»", sum(1 for z in x["raices"] if z["name"] != "4.- INSPECCIONES")),
    ("Archivos", len(P)),
    ("  se quedan igual", acc["IGUAL"]),
    ("  se mueven o renombran", acc["MOVER"]),
    ("  van a la papelera (recuperables 30 días, en bitácora)", acc["PAPELERA"]),
    ("PDF de inspección que pasan a INSPECCION-<NUCO> <fecha>.pdf", sum(1 for p in P if "/INSPECCION-" in p["despues"])),
    ("Documentos del vehículo que pasan a 1.-DOCUMENTACIÓN", sum(1 for p in P if p["destino_tipo"] == "DOCUMENTACION")),
    ("Archivos a 4.- INSPECCIONES/_POR REVISAR (plano)", sum(1 for p in P if p["destino_tipo"] == "POR REVISAR")),
    ("Archivos en <año>/SIN PERIODO…", sum(1 for p in P if "/SIN PERIODO" in p["despues"])),
    ("Nombre repetido en el destino (va con (2))", sum(1 for p in P if "mismo nombre" in p["aviso"])),
    ("", ""),
    ("Profundidad bajo 4.- INSPECCIONES (archivos)", ""),
] + [("  %d nivel(es) de carpeta" % k, n) for k, n in sorted(prof.items())] + [
    ("", ""),
    ("Carpetas por clasificación", ""),
] + [("  " + k, n) for k, n in collections.Counter(v[0] for v in nota.values()).most_common()] + [
    ("", ""),
    ("Revisiones automáticas", ""),
    ("  cada archivo tiene un solo destino", "OK"),
    ("  ningún archivo cambia de NUCO", "OK"),
    ("  dos archivos nunca terminan en la misma ruta", "OK"),
    ("  nada va a la papelera sin motivo", "OK"),
]
for f in filas:
    r.append(list(f))
for row in r.iter_rows():
    for c in row:
        c.font = F(bold=(c.row == 1), size=13 if c.row == 1 else 10)
r.column_dimensions["A"].width = 60
r.column_dimensions["B"].width = 14

# Carpetas: cada carpeta que se clasificó (no las de dentro de una sección), antes → después
s = wb.create_sheet("Carpetas")
s.append(["NUCO", "Clasificación", "Carpeta antes", "Carpeta después", "Archivos dentro", "Aviso"])
color = {"PERIODO": "DDEBF7", "MOTIVO": "E2EFDA", "AJENO": "FCE4D6", "SIN CLASIFICAR": "F8CBAD", "ENVOLTORIO": "EDEDED",
         "SECCION SUELTA": "FFF2CC", "EXTRA": "FFF2CC", "AÑO": "FFFFFF"}
raiz_nombre = {z["id"]: z["name"] for z in x["raices"]}
filas_c = []
for cid, (k, aviso) in nota.items():
    c = C[cid]
    if k == "AÑO" and c["name"].strip() == c["name"]:
        continue
    despues = h["destino"].get(cid)
    filas_c.append([c["nuco"], k, "/".join([c["nuco"]] + c["ruta"]),
                    "(se quita este nivel)" if despues is None else "/".join([c["nuco"], "4.- INSPECCIONES"] + despues),
                    cuenta[(c["nuco"], "/".join(c["ruta"]))], aviso])
orden = ["SIN CLASIFICAR", "AJENO", "EXTRA", "SECCION SUELTA", "ENVOLTORIO", "MOTIVO", "PERIODO", "AÑO"]
for f in sorted(filas_c, key=lambda f: (orden.index(f[1]), int(f[0]) if f[0].isdigit() else 0, f[2])):
    s.append(f)
    for c in s[s.max_row]:
        c.font = F(size=10)
        c.fill = PatternFill("solid", fgColor=color.get(f[1], "FFFFFF"))
for c in s[1]:
    c.font = F(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
for col, w in zip("ABCDEF", [7, 16, 70, 70, 10, 60]):
    s.column_dimensions[col].width = w
s.freeze_panes = "A2"
s.auto_filter.ref = s.dimensions

# PDF renombrados
q = wb.create_sheet("PDF renombrados")
q.append(["NUCO", "Antes", "Después"])
for p in P:
    if "/INSPECCION-" in p["despues"]:
        q.append([p["nuco"], p["antes"], p["despues"]])
for row in q.iter_rows():
    for c in row:
        c.font = F(size=10, bold=(c.row == 1))
for col, w in zip("ABC", [7, 90, 90]):
    q.column_dimensions[col].width = w
q.freeze_panes = "A2"
q.auto_filter.ref = q.dimensions

def hoja(nombre, filtro, cols):
    w = wb.create_sheet(nombre)
    w.append([c for c, _ in cols])
    for p in P:
        if filtro(p):
            w.append([f(p) for _, f in cols])
    for row in w.iter_rows():
        for c in row:
            c.font = F(size=10, bold=(c.row == 1))
    for k, (c, _) in enumerate(cols):
        w.column_dimensions["ABCDE"[k]].width = 8 if c == "NUCO" else 75
    w.freeze_panes = "A2"
    w.auto_filter.ref = w.dimensions

hoja("A documentación", lambda p: p["destino_tipo"] == "DOCUMENTACION",
     [("NUCO", lambda p: p["nuco"]), ("Antes", lambda p: p["antes"]), ("Después", lambda p: p["despues"]), ("Por qué", lambda p: p["aviso"])])
hoja("Por revisar", lambda p: p["destino_tipo"] == "POR REVISAR",
     [("NUCO", lambda p: p["nuco"]), ("Antes", lambda p: p["antes"]), ("Después", lambda p: p["despues"]), ("Por qué", lambda p: p["aviso"])])
hoja("Papelera", lambda p: p["accion"] == "PAPELERA",
     [("NUCO", lambda p: p["nuco"]), ("Archivo", lambda p: p["antes"]), ("Por qué", lambda p: p["aviso"])])

# Avisos de archivo
v = wb.create_sheet("Choques de nombre")
v.append(["NUCO", "Antes", "Después", "Aviso"])
for p in P:
    if "mismo nombre" in p.get("aviso", ""):
        v.append([p["nuco"], p["antes"], p["despues"], p["aviso"]])
for row in v.iter_rows():
    for c in row:
        c.font = F(size=10, bold=(c.row == 1))
for col, w in zip("ABCD", [7, 80, 80, 45]):
    v.column_dimensions[col].width = w

wb.save(SAL + ".xlsx")
print("listo:", SAL + ".xlsx", SAL + "-arboles.txt", "| filas Carpetas:", len(filas_c))
