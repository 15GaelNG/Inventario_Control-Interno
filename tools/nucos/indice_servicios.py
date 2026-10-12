"""Índice de servicios: una hoja con cada factura (CFDI) guardada en 2.-SERVICIOS de cada NUCO, sin mover nada en Drive.

  indice_servicios.py            Lee .cache/servicios/foto.json (foto.py <salida> SERVIC) y .cache/servicios/xml.json
                                 (las facturas ya leídas) y escribe la hoja "INDICE SERVICIOS NUCO" en
                                 PRUEBA DE NUCOS VEHICULARES (la crea o la reemplaza) y una copia en Excel en .cache/servicios.

Pestañas:
  FACTURAS   un renglón por XML: NUCO, placa, fecha, tipo (Ingreso / Pago / Egreso), taller, RFC, total, folio fiscal,
             servicio (la carpeta del evento), año/mes de la carpeta del área, ligas al XML, a su PDF y a la carpeta.
  POR NUCO   cuántos servicios, total (solo Ingreso − Egreso, cada folio fiscal una vez), primero y último, taller más usado.
  TALLERES   lo mismo por taller.

Reglas (12-oct-2026): los complementos de PAGO no suman (contarían doble); una nota de crédito (Egreso) resta; una factura
repetida (mismo UUID en dos carpetas) suma una vez. El vehículo es su NUCO (la carpeta), nunca el folio.
"""
import sys, json, re, collections, datetime
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e, reglas

C = e.CACHE / "servicios"
NOMBRE = "INDICE SERVICIOS NUCO"
TIPO = {"I": "Ingreso", "P": "Pago", "E": "Egreso", "T": "Traslado", "N": "Nómina"}
SUB = re.compile(r"^\d*\W*(FACTURAS?|COTIZACI\w*|AUTORIZACI\w*|GALERIA|EVIDENCIA|FOTOS|ENCUESTAS?)$")
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
CONTENEDOR = re.compile(r"LIBRETA|^SERVICIOS?$|SERVICIO VEHICULAR|^\d+\W*SERVICIOS?$")
liga = lambda i: "https://drive.google.com/file/d/%s/view" % i
carpeta_liga = lambda i: "https://drive.google.com/drive/folders/%s" % i
num = lambda s: str(int(re.match(r"\s*(\d+)", s).group(1))) if re.match(r"\s*(\d+)", s or "") else ""


def mes_de(nombre):
    t = reglas.normal(nombre)
    for i, m in enumerate(MESES):
        if re.search(r"\b" + m[:4], t):
            return i + 1
    return None


def filas():
    foto = json.load(open(C / "foto.json", encoding="utf-8"))
    xml = json.load(open(C / "xml.json", encoding="utf-8"))
    vehiculos = {x["nuco"]: x for x in json.load(open(e.CACHE / "verificaciones" / "vehiculos.json", encoding="utf-8"))}
    carpetas = {c["id"]: c for c in foto["carpetas"]}
    por_padre = collections.defaultdict(list)
    for a in foto["archivos"]:
        por_padre[a["padre"]].append(a)
    out = []
    for a in foto["archivos"]:
        x = xml.get(a["id"])
        if not x or a["nuco"].startswith("_"):
            continue
        r = [s.strip() for s in a["ruta"][1:-1]]          # entre 2.-SERVICIOS y el archivo
        # El evento: la carpeta que contiene FACTURA/…; si el XML está directo en el evento, esa misma
        i = len(r) - 1
        while i >= 0 and SUB.match(reglas.normal(r[i])):
            i -= 1
        evento = r[i] if i >= 0 and not CONTENEDOR.search(reglas.normal(r[i])) and not re.fullmatch(r"20\d\d", r[i]) and not mes_de(r[i]) else ""
        anio = next((s for s in r if re.fullmatch(r"20\d\d", s)), "")
        mes = next((mes_de(s) for s in r if mes_de(s) and len(s) < 16), None)
        # El evento: subir desde el archivo hasta la carpeta del evento (para su liga)
        pid, ev_id = a["padre"], None
        for _ in range(len(r) - i):
            if pid in carpetas:
                ev_id = pid
                pid = carpetas[pid]["padre"]
        base = re.sub(r"\.xml$", "", a["name"], flags=re.I).strip().upper()
        pdf = next((b for b in por_padre[a["padre"]] if b["name"].lower().endswith(".pdf") and re.sub(r"\.pdf$", "", b["name"], flags=re.I).strip().upper() == base), None)
        n = num(a["nuco"])
        try:
            total = float(x.get("total") or 0)
        except ValueError:
            total = 0.0
        out.append(dict(nuco=n, placa=vehiculos.get(n, {}).get("placa", ""), fecha=(x.get("fecha") or "")[:10],
                        tipo=TIPO.get(x.get("tipo"), x.get("tipo") or "?"), taller=x.get("emisor", ""), rfc=x.get("rfc", ""),
                        total=total, uuid=(x.get("uuid") or "").upper(), servicio=evento, anio_carpeta=anio,
                        mes_carpeta=MESES[mes - 1] if mes else "", xml=liga(a["id"]), pdf=liga(pdf["id"]) if pdf else "",
                        carpeta=carpeta_liga(ev_id or a["padre"]), ruta="/".join(a["ruta"])))
    out.sort(key=lambda f: (int(f["nuco"]) if f["nuco"] else 0, f["fecha"]))
    vistos = set()
    for f in out:
        f["cuenta"] = f["tipo"] in ("Ingreso", "Egreso") and f["uuid"] not in vistos
        if f["uuid"]:
            vistos.add(f["uuid"])
        f["repetida"] = "" if f["cuenta"] or f["tipo"] not in ("Ingreso", "Egreso") else "sí (mismo folio fiscal en otra carpeta)"
    return out


def resumen(out, clave):
    g = collections.defaultdict(lambda: dict(servicios=set(), total=0.0, fechas=[], talleres=collections.Counter(), facturas=0))
    for f in out:
        if not f["cuenta"]:
            continue
        x = g[clave(f)]
        x["facturas"] += 1
        x["total"] += f["total"] if f["tipo"] == "Ingreso" else -f["total"]
        x["servicios"].add(f["servicio"] or f["carpeta"])
        if f["fecha"]:
            x["fechas"].append(f["fecha"])
        x["talleres"][f["taller"]] += 1
    return g


def escribir():
    out = filas()
    enc = ["NUCO", "PLACA", "FECHA FACTURA", "TIPO", "TALLER", "RFC", "TOTAL", "CUENTA EN TOTALES", "REPETIDA", "FOLIO FISCAL (UUID)",
           "SERVICIO (CARPETA)", "AÑO CARPETA", "MES CARPETA", "XML", "PDF", "CARPETA DEL SERVICIO", "RUTA"]
    fac = [enc] + [[f["nuco"], f["placa"], f["fecha"], f["tipo"], f["taller"], f["rfc"], f["total"], "SÍ" if f["cuenta"] else "NO",
                    f["repetida"], f["uuid"], f["servicio"], f["anio_carpeta"], f["mes_carpeta"], f["xml"], f["pdf"], f["carpeta"], f["ruta"]] for f in out]
    g = resumen(out, lambda f: f["nuco"])
    pnuco = [["NUCO", "PLACA", "FACTURAS", "SERVICIOS", "TOTAL", "PRIMERA", "ÚLTIMA", "TALLER MÁS USADO"]] + [
        [n, next((f["placa"] for f in out if f["nuco"] == n), ""), x["facturas"], len(x["servicios"]), round(x["total"], 2),
         min(x["fechas"]) if x["fechas"] else "", max(x["fechas"]) if x["fechas"] else "", x["talleres"].most_common(1)[0][0]]
        for n, x in sorted(g.items(), key=lambda kv: int(kv[0]) if kv[0] else 0)]
    t = resumen(out, lambda f: f["taller"])
    ptall = [["TALLER", "FACTURAS", "SERVICIOS", "TOTAL", "PRIMERA", "ÚLTIMA"]] + [
        [k, x["facturas"], len(x["servicios"]), round(x["total"], 2), min(x["fechas"]) if x["fechas"] else "", max(x["fechas"]) if x["fechas"] else ""]
        for k, x in sorted(t.items(), key=lambda kv: -kv[1]["total"])]
    hojas = {"FACTURAS": fac, "POR NUCO": pnuco, "TALLERES": ptall, "LEEME": [
        ["Índice de servicios (facturas CFDI en 2.-SERVICIOS de cada NUCO). Solo lectura: no se movió nada en Drive."],
        ["Generado: " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M") + " con tools/nucos/indice_servicios.py"],
        ["TOTAL suma solo facturas de Ingreso menos notas de crédito (Egreso), cada folio fiscal una vez."],
        ["Los complementos de Pago no suman (son el pago de una factura que ya está). Montos con IVA (Total del CFDI)."],
        ["FECHA FACTURA es la del CFDI; AÑO/MES CARPETA, donde el área lo guardó (la carpeta 2024 trae servicios de 2022-2023)."]]}
    # Excel local
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    for nombre, datos in hojas.items():
        ws = wb.create_sheet(nombre)
        for fila in datos:
            ws.append(fila)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
    wb.save(C / "indice-servicios.xlsx")
    # La hoja de Google, en PRUEBA DE NUCOS VEHICULARES
    from autorizar import credenciales
    from googleapiclient.discovery import build
    d = e.drive()
    sh = build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()
    hay = e.llamar(d.files().list(q="'%s' in parents and name = '%s' and trashed = false" % (e.RESPALDO, NOMBRE),
                                  fields="files(id)"))["files"]
    if hay:
        sid = hay[0]["id"]
        props = e.llamar(sh.get(spreadsheetId=sid, fields="sheets.properties"))["sheets"]
        existentes = {p["properties"]["title"]: p["properties"]["sheetId"] for p in props}
        req = [{"addSheet": {"properties": {"title": n}}} for n in hojas if n not in existentes]
        if req:
            e.llamar(sh.batchUpdate(spreadsheetId=sid, body={"requests": req}))
        for n in hojas:
            e.llamar(sh.values().clear(spreadsheetId=sid, range="'%s'" % n))
    else:
        sid = e.llamar(d.files().create(body={"name": NOMBRE, "mimeType": "application/vnd.google-apps.spreadsheet",
                                              "parents": [e.RESPALDO]}, fields="id"))["id"]
        props = e.llamar(sh.get(spreadsheetId=sid, fields="sheets.properties"))["sheets"]
        req = [{"updateSheetProperties": {"properties": {"sheetId": props[0]["properties"]["sheetId"], "title": "FACTURAS"}, "fields": "title"}}]
        req += [{"addSheet": {"properties": {"title": n}}} for n in hojas if n != "FACTURAS"]
        e.llamar(sh.batchUpdate(spreadsheetId=sid, body={"requests": req}))
    e.llamar(sh.values().batchUpdate(spreadsheetId=sid, body={"valueInputOption": "RAW", "data": [
        {"range": "'%s'!A1" % n, "values": datos} for n, datos in hojas.items()]}))
    props = e.llamar(sh.get(spreadsheetId=sid, fields="sheets.properties"))["sheets"]
    e.llamar(sh.batchUpdate(spreadsheetId=sid, body={"requests": [
        {"updateSheetProperties": {"properties": {"sheetId": p["properties"]["sheetId"], "gridProperties": {"frozenRowCount": 1}},
                                   "fields": "gridProperties.frozenRowCount"}} for p in props] + [
        {"repeatCell": {"range": {"sheetId": p["properties"]["sheetId"], "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}}, "fields": "userEnteredFormat.textFormat.bold"}} for p in props]}))
    cuentan = [f for f in out if f["cuenta"]]
    print("facturas:", len(out), "| cuentan en totales:", len(cuentan), "| NUCO:", len(g), "| talleres:", len(t))
    print("total:", round(sum(f["total"] if f["tipo"] == "Ingreso" else -f["total"] for f in cuentan), 2))
    print("sin servicio identificado:", sum(1 for f in out if not f["servicio"]), "| con PDF:", sum(1 for f in out if f["pdf"]))
    print("hoja: https://docs.google.com/spreadsheets/d/%s/edit" % sid)


if __name__ == "__main__":
    escribir()
