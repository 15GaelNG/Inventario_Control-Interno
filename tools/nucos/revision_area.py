"""SOLO LECTURA: el Excel para que el área resuelva lo que quedó en duda en NUCOS VEHICULOS (15-oct-2026).

    revision_area.py <salida.xlsx>

Lee Drive en vivo, por niveles y de 40 en 40 carpetas (no carpeta por carpeta), y arma:
  RESUMEN                 cuánto hay en duda por carpeta del expediente
  INSPECCIONES SIN PERIODO  una inspección por renglón (4.-INSPECCIONES/<año>/SIN PERIODO - …): de qué periodo es
  SERVICIOS SIN FECHA     un servicio por renglón (2.-SERVICIOS/1.-LIBRETA DE SERVICIOS/SIN FECHA/<servicio>), con la
                          fecha que sugieren sus fotos (del plan de servicios)
  POR REVISAR             un archivo por renglón, de todos los _POR REVISAR (de cada carpeta y de la raíz de NUCOS)
Cada renglón trae su liga y columnas en blanco para la decisión del área.
"""
import sys, json, re, collections
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e, reglas

CARP = reglas.CARPETA
N = lambda s: re.sub(r"\s+", "", reglas.normal(s or ""))
liga_c = lambda i: "https://drive.google.com/drive/folders/" + i
liga_a = lambda i: "https://drive.google.com/file/d/%s/view" % i
RAICES = {"1.-DOCUMENTACION": "1.-DOCUMENTACIÓN", "2.-SERVICIOS": "2.-SERVICIOS", "3.-VERIFICACIONES": "3.-VERIFICACIONES",
          "4.-INSPECCIONES": "4.-INSPECCIONES"}


def hijos(d, ids):
    return e.hijos_de(d, list(ids)) if ids else {}


def contenido(d, tops):
    """{carpeta top: [archivos (con 'sub': ruta relativa)]} de todo lo que hay debajo, por niveles en lote."""
    out = {t: [] for t in tops}
    nivel = {t: (t, []) for t in tops}          # carpeta → (top, ruta relativa)
    while nivel:
        hs = hijos(d, nivel)
        sig = {}
        for cid, l in hs.items():
            top, ruta = nivel[cid]
            for h in l:
                if h["mimeType"] == CARP:
                    sig[h["id"]] = (top, ruta + [h["name"]])
                else:
                    out[top].append(dict(h, sub="/".join(ruta)))
        nivel = sig
    return out


def main(salida):
    d = e.drive()
    nucos = {f["id"]: f["name"] for f in hijos(d, [e.NUCOS])[e.NUCOS] if f["mimeType"] == CARP and f["name"].isdigit()}
    raiz_nucos = hijos(d, [e.NUCOS])[e.NUCOS]
    # Las 4 carpetas de cada NUCO
    raices = {}                                   # id → (nuco, carpeta)
    for nid, l in hijos(d, nucos).items():
        for h in l:
            k = next((v for k_, v in RAICES.items() if N(h["name"]).startswith(N(k_))), None)
            if h["mimeType"] == CARP and k:
                raices[h["id"]] = (nucos[nid], k)
    nivel1 = hijos(d, raices)
    # _POR REVISAR de cada carpeta; en inspecciones los años; en servicios la libreta
    por_revisar, anios, libretas = {}, {}, {}
    for rid, l in nivel1.items():
        nuco, carpeta = raices[rid]
        for h in l:
            if h["mimeType"] != CARP:
                continue
            if h["name"] == "_POR REVISAR":
                por_revisar[h["id"]] = (nuco, carpeta)
            elif carpeta == "4.-INSPECCIONES" and re.fullmatch(r"20\d\d", h["name"].strip()):
                anios[h["id"]] = (nuco, h["name"].strip())
            elif carpeta == "2.-SERVICIOS" and h["name"] == "1.-LIBRETA DE SERVICIOS":
                libretas[h["id"]] = nuco
    sin_periodo = {}
    for aid, l in hijos(d, anios).items():
        for h in l:
            if h["mimeType"] == CARP and h["name"].startswith("SIN PERIODO"):
                sin_periodo[h["id"]] = (anios[aid][0], anios[aid][1], h["name"])
    sin_fecha = {}
    for lid, l in hijos(d, libretas).items():
        for h in l:
            if h["mimeType"] == CARP and h["name"] == "SIN FECHA":
                sin_fecha[h["id"]] = libretas[lid]
    servicios = {}
    for sid, l in hijos(d, sin_fecha).items():
        for h in l:
            if h["mimeType"] == CARP:
                servicios[h["id"]] = (sin_fecha[sid], h["name"])
    raiz_pr = [h for h in raiz_nucos if h["name"] == "_POR REVISAR" and h["mimeType"] == CARP]
    for h in raiz_pr:
        por_revisar[h["id"]] = ("", "NUCOS (raíz)")
    print("SIN PERIODO %d · servicios SIN FECHA %d · _POR REVISAR %d" % (len(sin_periodo), len(servicios), len(por_revisar)), flush=True)
    cont = contenido(d, list(sin_periodo) + list(servicios) + list(por_revisar))

    # Fecha sugerida de cada servicio sin fecha (plan de servicios: por NUCO y nombre del servicio)
    sug = {}
    plan = e.CACHE / "servicios" / "20261015c" / "plan.json"
    if plan.exists():
        for v in json.load(open(plan, encoding="utf-8"))["eventos"].values():
            if not v["fecha"]:
                sug[(v["nuco"], v["nombre"])] = v["sugerencia"]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    res = wb.active
    res.title = "RESUMEN"
    ws_i = wb.create_sheet("INSPECCIONES SIN PERIODO")
    ws_s = wb.create_sheet("SERVICIOS SIN FECHA")
    ws_r = wb.create_sheet("POR REVISAR")
    ws_i.append(["NUCO", "Año (carpeta)", "Inspección (nombre que traía)", "Archivos", "PDF dentro", "Liga",
                 "Periodo correcto (área)", "Comentario"])
    for cid, (nuco, anio, nombre) in sorted(sin_periodo.items(), key=lambda x: (int(x[1][0]), x[1][1], x[1][2])):
        arch = cont[cid]
        ws_i.append([int(nuco), anio, nombre.replace("SIN PERIODO - ", "", 1), len(arch),
                     sum(1 for a in arch if a["name"].lower().endswith(".pdf")), liga_c(cid), "", ""])
    ws_s.append(["NUCO", "Servicio", "Archivos", "¿Trae factura?", "Fecha sugerida (por las fotos)", "Liga",
                 "Año y mes correctos (área)", "Comentario"])
    for cid, (nuco, nombre) in sorted(servicios.items(), key=lambda x: (int(x[1][0]), x[1][1])):
        arch = cont[cid]
        factura = any(re.search(r"FACTURA", N(a["sub"])) for a in arch)
        ws_s.append([int(nuco), nombre, len(arch), "sí (sin XML)" if factura else "no",
                     sug.get((nuco, nombre), ""), liga_c(cid), "", ""])
    ws_r.append(["NUCO", "Carpeta", "Archivo", "Subcarpeta", "Tipo", "Liga", "¿Qué es? / a dónde va (área)", "Comentario"])
    filas_r = []
    for cid, (nuco, carpeta) in por_revisar.items():
        for a in cont[cid]:
            filas_r.append([int(nuco) if nuco else "", carpeta, a["name"], a["sub"], a["mimeType"].split("/")[-1][:20],
                            liga_a(a["id"]), "", ""])
    for f in sorted(filas_r, key=lambda f: (f[0] == "", f[0] or 0, f[1], f[3], f[2])):
        ws_r.append(f)

    # Resumen
    res.append(["Qué", "Carpeta", "Unidades", "Archivos", "NUCO", "Pestaña"])
    def nu(ids, idx=0):
        return len({x[idx] for x in ids if x[idx]})
    res.append(["Inspecciones sin periodo", "4.-INSPECCIONES", len(sin_periodo), sum(len(cont[c]) for c in sin_periodo),
                nu(sin_periodo.values()), "INSPECCIONES SIN PERIODO"])
    res.append(["Servicios sin fecha", "2.-SERVICIOS", len(servicios), sum(len(cont[c]) for c in servicios),
                nu(servicios.values()), "SERVICIOS SIN FECHA"])
    por_c = collections.defaultdict(list)
    for cid, (nuco, carpeta) in por_revisar.items():
        por_c[carpeta].append((nuco, cid))
    for carpeta in ["4.-INSPECCIONES", "3.-VERIFICACIONES", "2.-SERVICIOS", "1.-DOCUMENTACIÓN", "NUCOS (raíz)"]:
        l = por_c.get(carpeta, [])
        res.append(["_POR REVISAR", carpeta, "", sum(len(cont[c]) for _, c in l), len({n for n, _ in l if n}), "POR REVISAR"])
    res.append([])
    res.append(["Cómo usarlo: cada pestaña trae su liga a Drive y dos columnas para el área (la decisión y un comentario)."])
    res.append(["Lo que el área decida se aplica después con bitácora; nadie mueve archivos a mano en NUCOS VEHICULOS."])
    res.append(["No incluye: los 84 certificados de verificación dudosos (en su lugar, Excel de verificaciones) ni lo apartado"])
    res.append(["(ANTERIORES / REPETIDOS / DUPLICADOS) ni _PAPELERA: no es duda."])

    cabeza, decide = PatternFill("solid", fgColor="1F3A5F"), PatternFill("solid", fgColor="FFF4CC")
    anchos = {"A": 8, "B": 26, "C": 48, "D": 14, "E": 22, "F": 52, "G": 30, "H": 30}
    for ws in (res, ws_i, ws_s, ws_r):
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cabeza
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        if ws is not res:
            ws.auto_filter.ref = ws.dimensions
            for col, w in anchos.items():
                ws.column_dimensions[col].width = w
            for fila in ws.iter_rows(min_row=2):
                for c in fila[6:8]:
                    c.fill = decide
                if isinstance(fila[5].value, str) and fila[5].value.startswith("http"):
                    fila[5].hyperlink = fila[5].value
                    fila[5].value = "Abrir"
                    fila[5].font = Font(color="1155CC", underline="single")
    for col, w in {"A": 26, "B": 20, "C": 10, "D": 10, "E": 8, "F": 28}.items():
        res.column_dimensions[col].width = w
    wb.save(salida)
    print("Excel:", salida)
    for r in res.iter_rows(min_row=2, max_row=8, values_only=True):
        print("  ", r)


if __name__ == "__main__":
    main(sys.argv[1])
