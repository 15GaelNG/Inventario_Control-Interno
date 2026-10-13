"""Inspecciones de la APP al expediente del NUCO (Fase 3, corrida 1).

    inspecciones_app.py plan <foto.json> <salida>          SOLO LECTURA → <salida>.json (el plan) y <salida>.xlsx
    inspecciones_app.py aplicar <plan.json> --nucos 1,2,3  EN LA REAL (o --todos). Bitácora .cache/bitacoras/<fecha>-insp-app.jsonl
    inspecciones_app.py verificar <bitácora> <plan.json>   Solo lectura: cada archivo donde dice y cada celda con su liga
    inspecciones_app.py deshacer <bitácora> [--nucos …]    Regresa las celdas (si nadie las cambió) y luego los archivos

<foto.json> es la de `foto.py <salida> INSPECC` (el árbol de 4.-INSPECCIONES de cada NUCO, con md5), tomada justo antes.

Por cada inspección de la hoja INSPECCION VEHICULAR de producción, con las MISMAS reglas que la app usa para lo nuevo
(InspeccionesService.registrar + ExpedienteNuco):

  PDF (columna FORMATO INSPECCION VEHICULAR, hoy una ruta "INSPECCIONES VEHICULARES/<tipo>/<archivo>.pdf" de REPORTES)
    - Si el área ya copió ese mismo PDF (mismo md5) a 4.-INSPECCIONES del NUCO: LIGA → la celda apunta a esa copia; el
      original de REPORTES se queda donde está (decisión del 9-oct).
    - Si no: MOVER → <NUCO>/4.-INSPECCIONES/<año>/<N> TRIMESTRE (antes de 2026: <N> BIMESTRE) con el nombre
      INSPECCION-<NUCO a 4> <aaaa-mm-dd>.pdf ((2), (3)… si ya existe), y la celda guarda su liga (mover conserva el ID).
  Firmas y diagramas marcados (INS FRONTAL/TRASERA/IZQUIERDA/DERECHA, FIRMA RESPONSABLE/INSPECTOR, hoy
  "INSPECCION VEHICULAR_Images/<archivo>")
    - MOVER → <la carpeta de su PDF>/FIRMAS Y DIAGRAMAS/<COLUMNA>-<NUCO a 4> <aaaa-mm-dd>.<ext>, y la celda guarda su liga.
    - El diagrama en blanco (MODELOS INSPECCION/…) es la plantilla: no se toca.

Nada se adivina: lo que no cuadra exacto (sin NUCO, sin fecha, el archivo no está o está dos veces, la cuenta no
puede moverlo, la celda trae otra cosa) sale como REVISAR y no se toca.
"""
import sys, json, re, collections
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e, reglas
from autorizar import credenciales
from googleapiclient.discovery import build

CARP = reglas.CARPETA
HOJA = "INSPECCION VEHICULAR"
COL_ID = "ID INSPECCION"
COL_PDF = "FORMATO INSPECCION VEHICULAR"
COLS_IMG = ["INS FRONTAL", "INS TRASERA", "INS IZQUIERDA", "INS DERECHA", "FIRMA RESPONSABLE", "FIRMA INSPECTOR"]
REPORTES = "1EYmA9RNLiPcY4vzZ07j2CC_0CBnutwkp"           # DRIVE_FOLDER_ID_REPORTES (prod): "INSPECCIONES VEHICULARES"
IMAGENES = "1X4je88ejXgBXpELZFYo2-pftsFjHZabL"           # DRIVE_FOLDER_ID_INSPECCIONES_IMAGENES (prod)
RAIZ_INSP = "4.-INSPECCIONES"
FIRMAS = "FIRMAS Y DIAGRAMAS"
ORDINAL = ["", "1ER", "2DO", "3ER", "4TO", "5TO", "6TO"]
CAMPOS = "id,name,mimeType,md5Checksum,parents,trashed,owners(emailAddress),capabilities(canMoveItemWithinDrive,canRename,canEdit)"
liga = lambda i: "https://drive.google.com/file/d/%s/view?usp=drivesdk" % i


def clave(nombre):
    """Como ExpedienteNuco.clave_: sin espacios ni acentos, en mayúsculas ("4.- INSPECCIONES" = "4.-INSPECCIONES")."""
    return re.sub(r"\s+", "", reglas.normal(nombre))


def letra(n):
    s, n = "", n + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def periodo(anio, mes):
    """ExpedienteNuco.periodoInspeccion: desde 2026 trimestre, antes bimestre."""
    return ORDINAL[(mes - 1) // 3 + 1] + " TRIMESTRE" if anio >= 2026 else ORDINAL[(mes - 1) // 2 + 1] + " BIMESTRE"


def fecha_de(texto):
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto or "")
    if not m:
        return None
    dia, mes, anio = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return (anio, mes, dia) if 1 <= mes <= 12 and 1 <= dia <= 31 else None


def listar(d, padre):
    """Todos los hijos de una carpeta, con dueño y permisos (para saber si la cuenta puede moverlos)."""
    out, tok = [], None
    while True:
        r = e.llamar(d.files().list(q="'%s' in parents and trashed = false" % padre, pageSize=1000, pageToken=tok,
                                    fields="nextPageToken,files(%s)" % CAMPOS, supportsAllDrives=True, includeItemsFromAllDrives=True))
        out += r.get("files", [])
        tok = r.get("nextPageToken")
        if not tok:
            return out


def plan(foto_json, salida):
    d = e.drive()
    s = build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()

    # ---- la hoja ----
    v = e.llamar(s.values().get(spreadsheetId=e.LIBRO, range="'%s'!A1:ZZ100000" % HOJA)).get("values", [])
    H = [x.strip() for x in v[0]]
    for c in [COL_ID, "NUCO", "FECHA", COL_PDF] + COLS_IMG:
        if c not in H:
            raise SystemExit("La hoja no tiene la columna %r" % c)
    filas = []
    for n, r in enumerate(v[1:], start=2):
        r = r + [""] * (len(H) - len(r))
        if any(str(x).strip() for x in r):
            filas.append(dict(fila=n, **{c: r[H.index(c)].strip() for c in [COL_ID, "NUCO", "FECHA", COL_PDF] + COLS_IMG}))
    celda = lambda col, fila: "%s%d" % (letra(H.index(col)), fila)

    # ---- los archivos de origen ----
    # REPORTES: "INSPECCIONES VEHICULARES/<subcarpeta>/<archivo>" (se resuelve la ruta exacta, nivel por nivel)
    rep = {"": listar(d, REPORTES)}
    for sub in [x for x in rep[""] if x["mimeType"] == CARP]:
        rep[sub["name"]] = listar(d, sub["id"])
    img = listar(d, IMAGENES)
    por_nombre_img = collections.defaultdict(list)
    for f in img:
        if f["mimeType"] != CARP:
            por_nombre_img[f["name"]].append(f)

    def resolver_pdf(ruta):
        partes = ruta.split("/")
        if len(partes) == 3 and partes[0] == "INSPECCIONES VEHICULARES":
            subs = [x for x in rep[""] if x["mimeType"] == CARP and x["name"] == partes[1]]
            if len(subs) != 1:
                return None, "%d carpetas %r en REPORTES" % (len(subs), partes[1])
            m = [x for x in rep[partes[1]] if x["name"] == partes[2] and x["mimeType"] != CARP]
        elif len(partes) == 2 and partes[0] == "INSPECCIONES VEHICULARES":
            m = [x for x in rep[""] if x["name"] == partes[1] and x["mimeType"] != CARP]
        else:
            return None, "ruta que no es de REPORTES: %r" % ruta
        if len(m) != 1:
            return None, "%d archivos con ese nombre en esa carpeta" % len(m)
        return m[0], ""

    def resolver_img(ruta):
        partes = ruta.split("/")
        if len(partes) != 2 or partes[0] != "INSPECCION VEHICULAR_Images":
            return None, "ruta que no es de INSPECCION VEHICULAR_Images: %r" % ruta
        m = por_nombre_img.get(partes[1], [])
        if len(m) != 1:
            return None, "%d archivos con ese nombre en INSPECCION VEHICULAR_Images" % len(m)
        return m[0], ""

    # ---- la foto de 4.-INSPECCIONES de cada NUCO ----
    foto = json.loads(Path(foto_json).read_text(encoding="utf-8"))
    nucos = {x["name"].strip(): x for x in foto["nucos"]}
    raices = collections.defaultdict(list)
    for r in foto["raices"]:
        raices[r["nuco"].strip()].append(r)
    hijos = collections.defaultdict(list)          # id de carpeta → [carpetas y archivos dentro]
    por_id = {}
    for x in foto["carpetas"]:
        hijos[x["padre"]].append(dict(x, carpeta=True))
        por_id[x["id"]] = x
    pdfs_por_md5 = collections.defaultdict(list)    # (nuco, md5) → archivos
    for x in foto["archivos"]:
        hijos[x["padre"]].append(dict(x, carpeta=False))
        if x.get("md5"):
            pdfs_por_md5[(x["nuco"].strip(), x["md5"])].append(x)

    # Carpetas que el plan va a crear: se identifican por (padre, nombre) para que dos inspecciones usen la misma
    nuevas = {}
    nombres_tomados = collections.defaultdict(set)   # carpeta (id o "nueva:…") → nombres que ya hay o que el plan pone

    def tomados(carpeta):
        if carpeta not in nombres_tomados and not carpeta.startswith("nueva:"):
            nombres_tomados[carpeta] = {x["name"] for x in hijos.get(carpeta, [])}
        return nombres_tomados[carpeta]

    def subcarpeta(padre, nombre):
        """La carpeta `nombre` dentro de `padre` (como ExpedienteNuco.subcarpeta_: exacta, o igual sin espacios ni
        acentos). Si no existe, una nueva del plan. Si hay dos que coinciden, no se escoge: error."""
        if not padre.startswith("nueva:"):
            dentro = [x for x in hijos.get(padre, []) if x["carpeta"]]
            exacta = [x for x in dentro if x["name"] == nombre]
            if len(exacta) == 1:
                return exacta[0]["id"], ""
            parecidas = [x for x in dentro if clave(x["name"]) == clave(nombre)]
            if len(parecidas) == 1:
                return parecidas[0]["id"], ""
            if len(exacta) + len(parecidas) > 1:
                return None, "hay %d carpetas %r en %s" % (len(exacta) + len(parecidas), nombre, ruta_de(padre))
        k = (padre, nombre)
        if k not in nuevas:
            nuevas[k] = "nueva:%d" % len(nuevas)
            nombres_tomados[nuevas[k]] = set()
        return nuevas[k], ""

    def ruta_de(carpeta):
        if carpeta.startswith("nueva:"):
            padre, nombre = next(k for k, v in nuevas.items() if v == carpeta)
            return ruta_de(padre) + "/" + nombre
        if carpeta in por_id:
            return "/".join([por_id[carpeta]["nuco"]] + por_id[carpeta]["ruta"])
        for rs in raices.values():
            for r in rs:
                if r["id"] == carpeta:
                    return r["nuco"] + "/" + r["name"]
        return carpeta

    def nombre_libre(carpeta, prefijo, nuco, fecha, ext):
        """ExpedienteNuco.nombreConFecha: PREFIJO-0088 2026-05-12.ext, y (2), (3)… si ya existe."""
        base = "%s-%04d %04d-%02d-%02d" % (prefijo, int(nuco), fecha[0], fecha[1], fecha[2])
        ya = tomados(carpeta)
        nombre, i = base + ext, 2
        while nombre in ya:
            nombre, i = "%s (%d)%s" % (base, i, ext), i + 1
        ya.add(nombre)
        return nombre

    # ---- el plan ----
    ops, revisar = [], []
    usados = collections.Counter()
    for f in filas:
        for c in [COL_PDF] + COLS_IMG:
            if f[c] and not f[c].startswith("http") and not f[c].startswith("MODELOS INSPECCION/"):
                usados[f[c]] += 1

    def anotar_revisar(f, col, motivo):
        revisar.append(dict(fila=f["fila"], id=f[COL_ID], nuco=f["NUCO"], fecha=f["FECHA"], columna=col,
                            celda=celda(col, f["fila"]) if col else "", valor=f.get(col, "") if col else "", motivo=motivo))

    for f in filas:
        if not f[COL_ID]:
            anotar_revisar(f, "", "renglón sin ID INSPECCION")
            continue
        pendientes = [c for c in [COL_PDF] + COLS_IMG if f[c] and not f[c].startswith("http") and not f[c].startswith("MODELOS INSPECCION/")]
        if not pendientes:
            continue   # todo es liga, plantilla o vacío: nada que hacer
        nuco = f["NUCO"].lstrip("0") if f["NUCO"].isdigit() else ""
        fecha = fecha_de(f["FECHA"])
        if not nuco or nuco not in nucos:
            for c in pendientes:
                anotar_revisar(f, c, "NUCO %r sin carpeta en NUCOS VEHICULOS" % f["NUCO"])
            continue
        if not fecha:
            for c in pendientes:
                anotar_revisar(f, c, "fecha ilegible: %r" % f["FECHA"])
            continue
        rs = raices.get(nuco, [])
        if len(rs) > 1:
            for c in pendientes:
                anotar_revisar(f, c, "el NUCO tiene %d carpetas de inspecciones" % len(rs))
            continue
        raiz = rs[0]["id"] if rs else subcarpeta(nucos[nuco]["id"], RAIZ_INSP)[0]
        per = periodo(fecha[0], fecha[1])
        anio_c, err = subcarpeta(raiz, str(fecha[0]))
        per_c, err2 = (subcarpeta(anio_c, per) if anio_c else (None, err))
        if not per_c:
            for c in pendientes:
                anotar_revisar(f, c, err2 or err)
            continue
        destino_insp = per_c    # dónde quedan sus firmas: junto a su PDF

        # El PDF
        if COL_PDF in pendientes:
            if f[COL_PDF].startswith("data:"):
                anotar_revisar(f, COL_PDF, "la celda trae %r, no una ruta" % f[COL_PDF][:20])
            elif usados[f[COL_PDF]] > 1:
                anotar_revisar(f, COL_PDF, "ese PDF lo citan %d renglones" % usados[f[COL_PDF]])
            else:
                src, motivo = resolver_pdf(f[COL_PDF])
                if not src:
                    anotar_revisar(f, COL_PDF, motivo)
                else:
                    copias = pdfs_por_md5.get((nuco, src.get("md5Checksum")), []) if src.get("md5Checksum") else []
                    if copias:
                        # La que esté en el periodo de la inspección, si hay; si no, la de ruta más corta
                        copias = sorted(copias, key=lambda x: (x["padre"] != per_c, len(x["ruta"])))
                        x = copias[0]
                        aviso = []
                        # Nunca adivina (13-oct, NUCO 320): si la copia del área dice OTRA fecha (en su nombre) o está
                        # en OTRO periodo, no se sabe cuál está mal; se revisa. Una carpeta de motivo (POR BAJA …) sí vale.
                        fecha_copia = re.search(r"INSPECCION-\d+ (\d{4})-(\d{2})-(\d{2})", x["name"])
                        carpeta_copia = x["ruta"][-2] if len(x["ruta"]) >= 2 else ""
                        es_periodo = re.search(r"(TRIMESTRE|BIMESTRE)$", reglas.normal(carpeta_copia))
                        motivo = ""
                        if fecha_copia and tuple(int(g) for g in fecha_copia.groups()) != fecha:
                            motivo = "la copia idéntica del área (%s) dice otra fecha que la inspección (%s)" % ("/".join(x["ruta"]), f["FECHA"][:10])
                        elif x["padre"] != per_c and es_periodo:
                            motivo = "la copia idéntica del área está en %s y la inspección es de %s (%s)" % ("/".join(x["ruta"][:-1]), f["FECHA"][:10], per)
                        if motivo:
                            # Sus firmas van junto a su PDF: tampoco se mueven hasta que se aclare
                            for c in pendientes:
                                anotar_revisar(f, c, motivo if c == COL_PDF else "va junto a su PDF, que está por revisar")
                            continue
                        if x["padre"] != per_c:
                            aviso.append("la copia del área está en %s, no en %s" % ("/".join(x["ruta"][:-1]), per))
                        if len(copias) > 1:
                            aviso.append("%d copias iguales en el NUCO; se usa %s" % (len(copias), "/".join(x["ruta"])))
                        destino_insp = x["padre"]
                        ops.append(dict(fila=f["fila"], id=f[COL_ID], nuco=nuco, fecha=f["FECHA"], columna=COL_PDF,
                                        celda=celda(COL_PDF, f["fila"]), accion="LIGA", valor_antes=f[COL_PDF],
                                        archivo=x["id"], origen=f[COL_PDF], destino="/".join([nuco] + x["ruta"]),
                                        valor_nuevo=liga(x["id"]), aviso="; ".join(aviso)))
                    elif not (src.get("capabilities") or {}).get("canMoveItemWithinDrive"):
                        anotar_revisar(f, COL_PDF, "la cuenta no puede mover el PDF (dueño %s)" %
                                       ",".join(o.get("emailAddress", "?") for o in src.get("owners", [])))
                    else:
                        nombre = nombre_libre(per_c, "INSPECCION", nuco, fecha, ".pdf")
                        ops.append(dict(fila=f["fila"], id=f[COL_ID], nuco=nuco, fecha=f["FECHA"], columna=COL_PDF,
                                        celda=celda(COL_PDF, f["fila"]), accion="MOVER", valor_antes=f[COL_PDF],
                                        archivo=src["id"], de_padre=(src.get("parents") or [""])[0], de_nombre=src["name"],
                                        a_padre=per_c, a_nombre=nombre, origen=f[COL_PDF],
                                        destino=ruta_de(per_c) + "/" + nombre, valor_nuevo=liga(src["id"]),
                                        aviso="" if per_c.startswith("nueva:") or not [n for n in tomados(per_c) if
                                              n.startswith("INSPECCION-%04d %04d-%02d-%02d" % (int(nuco), *fecha)) and n != nombre]
                                        else "ya había un PDF de inspección de ese día en la carpeta (el del área): quedan los dos"))

        # Firmas y diagramas
        firmas_c = None
        for c in [x for x in COLS_IMG if x in pendientes]:
            if f[c].startswith("data:"):
                anotar_revisar(f, c, "la celda trae %r, no una ruta" % f[c][:20])
                continue
            if usados[f[c]] > 1:
                anotar_revisar(f, c, "esa imagen la citan %d renglones" % usados[f[c]])
                continue
            src, motivo = resolver_img(f[c])
            if not src:
                anotar_revisar(f, c, motivo)
                continue
            if not (src.get("capabilities") or {}).get("canMoveItemWithinDrive"):
                anotar_revisar(f, c, "la cuenta no puede moverla (dueño %s)" % ",".join(o.get("emailAddress", "?") for o in src.get("owners", [])))
                continue
            if firmas_c is None:
                firmas_c, err = subcarpeta(destino_insp, FIRMAS)
                if not firmas_c:
                    anotar_revisar(f, c, err)
                    continue
            ext = Path(src["name"]).suffix.lower() or ".png"
            nombre = nombre_libre(firmas_c, c, nuco, fecha, ext)
            ops.append(dict(fila=f["fila"], id=f[COL_ID], nuco=nuco, fecha=f["FECHA"], columna=c, celda=celda(c, f["fila"]),
                            accion="MOVER", valor_antes=f[c], archivo=src["id"], de_padre=(src.get("parents") or [""])[0],
                            de_nombre=src["name"], a_padre=firmas_c, a_nombre=nombre, origen=f[c],
                            destino=ruta_de(firmas_c) + "/" + nombre, valor_nuevo=liga(src["id"]), aviso=""))

    # Las carpetas a crear, en orden (un padre antes que sus hijas)
    crear = [dict(clave=v, padre=k[0], nombre=k[1], ruta=ruta_de(v)) for k, v in sorted(nuevas.items(), key=lambda kv: int(kv[1].split(":")[1]))]
    # Ninguna operación repite archivo ni destino
    assert len({o["archivo"] for o in ops if o["accion"] == "MOVER"}) == sum(1 for o in ops if o["accion"] == "MOVER"), "un archivo se movería dos veces"
    assert len({(o["a_padre"], o["a_nombre"]) for o in ops if o["accion"] == "MOVER"}) == sum(1 for o in ops if o["accion"] == "MOVER"), "dos archivos al mismo nombre"
    assert len({o["celda"] for o in ops}) == len(ops), "una celda se cambiaría dos veces"

    out = dict(hoja=HOJA, libro=e.LIBRO, foto=str(foto_json), crear=crear, ops=ops, revisar=revisar)
    Path(salida + ".json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    excel(out, salida + ".xlsx", len(filas))
    cuenta = collections.Counter((o["accion"], "PDF" if o["columna"] == COL_PDF else "imagen") for o in ops)
    print("renglones:", len(filas), "| operaciones:", dict(cuenta), "| carpetas a crear:", len(crear), "| a revisar:", len(revisar))


def excel(p, archivo, renglones):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    r = wb.active
    r.title = "Resumen"
    ops = p["ops"]
    c = collections.Counter((o["accion"], "PDF" if o["columna"] == COL_PDF else "firma/diagrama") for o in ops)
    filas = [("Ensayo: inspecciones de la app → expediente del NUCO (Fase 3, corrida 1)", ""),
             ("Solo lectura: no se movió nada ni se cambió ninguna celda.", ""), ("", ""),
             ("Renglones de la hoja", renglones),
             ("PDF: solo cambia la liga (el área ya lo había copiado al NUCO)", c[("LIGA", "PDF")]),
             ("PDF: se mueve de REPORTES al NUCO", c[("MOVER", "PDF")]),
             ("Firmas y diagramas: se mueven a FIRMAS Y DIAGRAMAS", c[("MOVER", "firma/diagrama")]),
             ("Celdas que cambian (la liga nueva)", len(ops)),
             ("Carpetas a crear", len(p["crear"])),
             ("Con aviso", sum(1 for o in ops if o["aviso"])),
             ("A revisar (no se tocan)", len(p["revisar"]))]
    for x in filas:
        r.append(list(x))
    r["A1"].font = Font(name="Arial", bold=True, size=13)
    for row in r.iter_rows(min_row=2):
        for cel in row:
            cel.font = Font(name="Arial")
    r.column_dimensions["A"].width = 70

    def hoja(titulo, cols, datos, color=None):
        h = wb.create_sheet(titulo)
        h.append([t for t, _ in cols])
        for x in datos:
            h.append([x.get(k, "") for _, k in cols])
            fill = PatternFill("solid", fgColor=(color(x) if color else "FFFFFF"))
            for cel in h[h.max_row]:
                cel.font = Font(name="Arial", size=10)
                cel.fill = fill
        for cel in h[1]:
            cel.font = Font(name="Arial", bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="305496")
            cel.alignment = Alignment(wrap_text=True, vertical="center")
        for i in range(len(cols)):
            h.column_dimensions[letra(i)].width = 18 if i < 6 else 48
        h.freeze_panes = "A2"
        h.auto_filter.ref = h.dimensions
        return h

    hoja("Plan", [("Fila", "fila"), ("ID inspección", "id"), ("NUCO", "nuco"), ("Fecha", "fecha"), ("Columna", "columna"),
                  ("Celda", "celda"), ("Acción", "accion"), ("Hoy en la celda", "valor_antes"), ("Queda en", "destino"),
                  ("Liga nueva", "valor_nuevo"), ("Aviso", "aviso")],
         sorted(ops, key=lambda o: (int(o["nuco"]), o["fila"], o["columna"] != COL_PDF)),
         lambda o: "FFF2CC" if o["aviso"] else ("E2EFDA" if o["accion"] == "LIGA" else "DDEBF7"))
    hoja("Carpetas a crear", [("Ruta", "ruta")], p["crear"])
    hoja("Revisar", [("Fila", "fila"), ("ID inspección", "id"), ("NUCO", "nuco"), ("Fecha", "fecha"), ("Columna", "columna"),
                     ("Celda", "celda"), ("Hoy en la celda", "valor"), ("Por qué no se toca", "motivo")],
         p["revisar"], lambda o: "FCE4D6")
    wb.save(archivo)


# ======================================================================================================== en la real

def hoja_viva(s):
    """La hoja como está AHORA: {ID INSPECCION: (fila, renglón)} y los encabezados. Un ID repetido no se toca."""
    v = e.llamar(s.values().get(spreadsheetId=e.LIBRO, range="'%s'!A1:ZZ100000" % HOJA)).get("values", [])
    H = [x.strip() for x in v[0]]
    por_id, repetidos = {}, set()
    for n, r in enumerate(v[1:], start=2):
        r = r + [""] * (len(H) - len(r))
        i = r[H.index(COL_ID)].strip()
        if not i:
            continue
        if i in por_id:
            repetidos.add(i)
        por_id[i] = (n, r)
    for i in repetidos:
        por_id.pop(i, None)
    return H, por_id


def aplicar(plan_json, nucos=None):
    """NUCO por NUCO: revisa que todo siga como en el plan, crea las carpetas, mueve y renombra, y escribe las ligas.
    Lo que cambió desde el plan se salta (no se rompe). Cada paso a la bitácora, en el momento."""
    import time
    p = json.loads(Path(plan_json).read_text(encoding="utf-8"))
    d = e.drive()
    s = build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()
    ops = [o for o in p["ops"] if nucos is None or o["nuco"] in nucos]
    if not ops:
        raise SystemExit("Nada que hacer para esos NUCO")
    crear = {c["clave"]: c for c in p["crear"]}
    bit = e.Bitacora(e.CACHE / "bitacoras" / (time.strftime("%Y%m%d-%H%M%S") + "-insp-app.jsonl"))
    bit.anotar(accion="inicio", plan=str(plan_json), nucos=sorted({o["nuco"] for o in ops}, key=int))
    resumen = collections.Counter()
    creadas = {}                                  # "nueva:k" → id real

    def real(carpeta):
        """El id de una carpeta del plan; las nuevas se crean al usarlas (sus padres primero)."""
        if not carpeta.startswith("nueva:"):
            return carpeta
        if carpeta in creadas:
            return creadas[carpeta]
        c = crear[carpeta]
        padre = real(c["padre"])
        # Si alguien la creó mientras tanto (mismo nombre exacto), se usa esa; si hay dos, no se escoge
        ya = e.llamar(d.files().list(q="'%s' in parents and name = '%s' and mimeType = '%s' and trashed = false" %
                                     (padre, c["nombre"].replace("'", "\\'"), CARP), fields="files(id)",
                                     supportsAllDrives=True, includeItemsFromAllDrives=True))["files"]
        if len(ya) > 1:
            raise RuntimeError("hay %d carpetas %r en %s" % (len(ya), c["nombre"], c["ruta"]))
        if ya:
            creadas[carpeta] = ya[0]["id"]
        else:
            creadas[carpeta] = e.llamar(d.files().create(body={"name": c["nombre"], "mimeType": CARP, "parents": [padre]},
                                                         fields="id", supportsAllDrives=True))["id"]
            nuco = c["ruta"].split("/")[0]
            bit.anotar(accion="crear_carpeta", nuco=nuco, id=creadas[carpeta], padre=padre, nombre=c["nombre"], ruta=c["ruta"])
            resumen["carpetas creadas"] += 1
        return creadas[carpeta]

    H, por_id = hoja_viva(s)
    for nuco in sorted({o["nuco"] for o in ops}, key=int):
        print("== NUCO %s" % nuco)
        celdas = []
        for o in [x for x in ops if x["nuco"] == nuco]:
            def saltar(motivo):
                bit.anotar(accion="saltado", nuco=nuco, id_inspeccion=o["id"], columna=o["columna"], archivo=o["archivo"], motivo=motivo)
                resumen["saltados"] += 1
                print("   saltado %s %s: %s" % (o["id"], o["columna"], motivo))
            # La celda: se busca por ID INSPECCION (las filas se recorren si alguien borra un renglón) y debe seguir igual
            fila = por_id.get(o["id"])
            if not fila:
                saltar("el ID INSPECCION ya no está (o está repetido) en la hoja"); continue
            n, r = fila
            if (r[H.index("NUCO")].strip().lstrip("0")) != nuco:
                saltar("el renglón ya no es de ese NUCO"); continue
            if r[H.index(o["columna"])].strip() != o["valor_antes"]:
                saltar("la celda cambió desde el plan"); continue
            # El archivo: debe seguir donde y como dice el plan
            try:
                f = e.llamar(d.files().get(fileId=o["archivo"], fields="id,name,parents,trashed", supportsAllDrives=True))
            except Exception as err:   # noqa: BLE001
                saltar("el archivo no se encontró: %s" % err); continue
            if f.get("trashed"):
                saltar("el archivo está en la papelera"); continue
            if o["accion"] == "MOVER":
                if (f.get("parents") or [None])[0] != o["de_padre"] or f["name"] != o["de_nombre"]:
                    saltar("el archivo ya no está donde lo vio el plan"); continue
                try:
                    destino = real(o["a_padre"])
                except Exception as err:   # noqa: BLE001
                    saltar("carpeta de destino: %s" % err); continue
                choca = e.llamar(d.files().list(q="'%s' in parents and name = '%s' and trashed = false" % (destino, o["a_nombre"].replace("'", "\\'")),
                                                fields="files(id)", supportsAllDrives=True, includeItemsFromAllDrives=True))["files"]
                if choca:
                    saltar("ya hay un archivo %r en el destino" % o["a_nombre"]); continue
                e.llamar(d.files().update(fileId=o["archivo"], addParents=destino, removeParents=o["de_padre"], fields="id",
                                          supportsAllDrives=True, body={"name": o["a_nombre"],
                                          "appProperties": {"origen_padre": o["de_padre"], "origen_nombre": o["de_nombre"]}}))
                bit.anotar(accion="mover", nuco=nuco, id=o["archivo"], de_padre=o["de_padre"], de_nombre=o["de_nombre"],
                           a_padre=destino, a_nombre=o["a_nombre"], app=True, id_inspeccion=o["id"], columna=o["columna"])
                resumen["archivos movidos"] += 1
            celdas.append((o, n))
        # Las ligas de este NUCO, de una sola vez (justo después de mover sus archivos)
        if celdas:
            e.llamar(s.values().batchUpdate(spreadsheetId=e.LIBRO, body={"valueInputOption": "RAW", "data": [
                {"range": "'%s'!%s%d" % (HOJA, letra(H.index(o["columna"])), n), "values": [[o["valor_nuevo"]]]} for o, n in celdas]}))
            for o, n in celdas:
                bit.anotar(accion="celda", nuco=nuco, id_inspeccion=o["id"], columna=o["columna"], antes=o["valor_antes"],
                           despues=o["valor_nuevo"], liga=o["accion"] == "LIGA")
            resumen["celdas"] += len(celdas)
        print("   %d celdas" % len(celdas))
    bit.anotar(accion="fin", resumen=dict(resumen))
    print("\nResumen:", dict(resumen))
    print("Bitácora:", bit.ruta)
    e.respaldar_bitacora(bit.ruta)


def verificar(bitacora, plan_json):
    """Solo lectura: lo que dice la bitácora está en Drive y en la hoja, y cada operación del plan para esos NUCO se hizo
    (o se saltó con su motivo)."""
    d = e.drive()
    s = build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()
    lineas = [json.loads(x) for x in Path(bitacora).read_text(encoding="utf-8").splitlines() if x.strip()]
    nucos = set(next(x for x in lineas if x["accion"] == "inicio")["nucos"])
    p = json.loads(Path(plan_json).read_text(encoding="utf-8"))
    k, mal = collections.Counter(), []
    for x in lineas:
        if x["accion"] == "mover":
            f = e.llamar(d.files().get(fileId=x["id"], fields="name,parents,trashed", supportsAllDrives=True))
            ok = not f.get("trashed") and (f.get("parents") or [None])[0] == x["a_padre"] and f["name"] == x["a_nombre"]
            k["archivo en su lugar" if ok else "ARCHIVO DISTINTO"] += 1
            if not ok:
                mal.append((x["id"], x["a_nombre"], f))
        elif x["accion"] == "crear_carpeta":
            f = e.llamar(d.files().get(fileId=x["id"], fields="name,parents,trashed", supportsAllDrives=True))
            k["carpeta creada ok" if not f.get("trashed") and x["padre"] in (f.get("parents") or []) else "CARPETA DISTINTA"] += 1
    H, por_id = hoja_viva(s)
    for x in lineas:
        if x["accion"] == "celda":
            fila = por_id.get(x["id_inspeccion"])
            ok = bool(fila) and fila[1][H.index(x["columna"])].strip() == x["despues"]
            k["celda con su liga" if ok else "CELDA DISTINTA"] += 1
            if not ok:
                mal.append((x["id_inspeccion"], x["columna"], fila and fila[1][H.index(x["columna"])]))
    hechas = {(x["id_inspeccion"], x["columna"]) for x in lineas if x["accion"] in ("celda", "saltado")}
    faltan = [(o["id"], o["columna"]) for o in p["ops"] if o["nuco"] in nucos and (o["id"], o["columna"]) not in hechas]
    k["del plan sin hacer ni saltar"] = len(faltan)
    print(dict(k))
    for m in mal[:20]:
        print("  ", m)
    for f in faltan[:20]:
        print("   sin hacer:", f)


def deshacer(bitacora, nucos=None):
    """Primero las celdas (solo las que siguen con la liga que se puso: si alguien la cambió después, se respeta),
    luego los archivos y las carpetas creadas, con expediente.deshacer (de la última línea a la primera)."""
    s = build("sheets", "v4", credentials=credenciales("lab"), cache_discovery=False).spreadsheets()
    lineas = [json.loads(x) for x in Path(bitacora).read_text(encoding="utf-8").splitlines() if x.strip()]
    marca = Path(str(bitacora) + ".deshechos")
    ya = set(marca.read_text(encoding="utf-8").split()) if marca.exists() else set()
    H, por_id = hoja_viva(s)
    data, k = [], collections.Counter()
    for x in lineas:
        if x["accion"] != "celda" or (nucos and x["nuco"] not in nucos) or x["nuco"] in ya:
            continue
        fila = por_id.get(x["id_inspeccion"])
        if not fila or fila[1][H.index(x["columna"])].strip() != x["despues"]:
            k["celda que cambió después: no se toca"] += 1
            continue
        data.append({"range": "'%s'!%s%d" % (HOJA, letra(H.index(x["columna"])), fila[0]), "values": [[x["antes"]]]})
        k["celdas regresadas"] += 1
    for i in range(0, len(data), 200):
        e.llamar(s.values().batchUpdate(spreadsheetId=e.LIBRO, body={"valueInputOption": "RAW", "data": data[i:i + 200]}))
    print(dict(k))
    e.deshacer(bitacora, set(nucos) if nucos else None)


if __name__ == "__main__":
    a = sys.argv[1:]
    nucos_ = a[a.index("--nucos") + 1].split(",") if "--nucos" in a else None
    if len(a) == 3 and a[0] == "plan":
        plan(a[1], a[2])
    elif a and a[0] == "aplicar" and (nucos_ or "--todos" in a):
        aplicar(a[1], nucos_)
    elif a and a[0] == "verificar" and len(a) == 3:
        verificar(a[1], a[2])
    elif a and a[0] == "deshacer":
        deshacer(a[1], nucos_)
    else:
        print(__doc__)
