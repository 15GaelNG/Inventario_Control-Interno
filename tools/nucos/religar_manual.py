"""Re-ligar las verificaciones que el usuario revisó a mano en verificaciones-completo.xlsx (pestaña "Por revisar").

    religar_manual.py plan <xlsx>         SOLO LECTURA → .cache/verificaciones/religar/manual.json y manual.xlsx
    religar_manual.py aplicar             EN LA REAL. Bitácora .cache/bitacoras/<fecha>-religar-manual.jsonl (+ Drive)
    religar_manual.py verificar <bit>     Solo lectura
    religar_manual.py deshacer <bit>      Regresa cada celda (si nadie la cambió después)

Qué NUCO manda (decisión del usuario, 14-oct-2026):
  - NUCO CORRECTO, si lo escribió (o lo dejó como venía en los de confianza ALTA);
  - confianza MEDIA sin NUCO CORRECTO → el de la columna "NUCO según el certificado" (serie o placa leídas);
  - sin confianza y sin NUCO CORRECTO → duda: no se toca.
La fecha: FECHA CORRECTA si la escribió; si no, se queda la capturada. Una fecha leída con OCR nunca corrige sola.

Si el vehículo cambia se escriben ID VEHICULO, FOLIO VEHICULO, PLACA (la copia del vehículo correcto, como
religar_verif.py) y NUCCO. El vehículo se busca por NUCO, jamás por folio. Antes de escribir, cada renglón (buscado por
su ID) debe seguir con el NUCO que tenía en el Excel; si no, se salta.
"""
import sys, json, re, time, datetime, collections
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e
from religar_verif import hojas, leer, letra, num

D = e.CACHE / "verificaciones" / "religar"
PLAN = D / "manual.json"
HOJA = "VERIFICACIONES"
COLS = ["ID VEHICULO", "FOLIO VEHICULO", "PLACA", "NUCCO", "FECHA VERIFICACION"]
FECHA = "FECHA VERIFICACION"      # se escribe como fecha (USER_ENTERED, libro es_MX); lo demás tal cual (RAW)


def decisiones(xlsx):
    from openpyxl import load_workbook
    h = load_workbook(xlsx, data_only=True)["Por revisar"]
    enc = [c.value for c in h[1]]
    out = []
    for r in h.iter_rows(min_row=2):
        x = dict(zip(enc, [c.value for c in r]))
        if not x.get("ID"):
            continue
        z, t, conf = x.get("NUCO CORRECTO"), x.get("NUCO según el certificado"), x.get("Confianza") or ""
        if z not in (None, ""):
            nuco, fuente = num(str(z)), "NUCO CORRECTO" + (" (ALTA, como venía)" if conf == "ALTA" and str(z) == str(t) else "")
        elif conf == "MEDIA" and t not in (None, ""):
            nuco, fuente = num(str(t)), "MEDIA: la columna T"
        else:
            nuco, fuente = "", "duda"
        fc = x.get("FECHA CORRECTA")
        if isinstance(fc, datetime.datetime):
            fc = fc.strftime("%d/%m/%Y")
        elif fc:
            m = re.match(r"\s*(\d{4})-(\d{2})-(\d{2})|\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(fc))
            if not m:
                raise SystemExit("FECHA CORRECTA ilegible en %s: %r" % (x["ID"], fc))
            fc = "%02d/%02d/%s" % ((int(m.group(3)), int(m.group(2)), m.group(1)) if m.group(1) else (int(m.group(4)), int(m.group(5)), m.group(6)))
        out.append(dict(id=x["ID"], nuco_excel=num(str(x.get("NUCO hoy") or "")), nuco=nuco, fuente=fuente, fecha=fc or "", confianza=conf))
    return out


def plan(xlsx):
    h = hojas()
    enc, ver = leer(h, HOJA)
    _, veh = leer(h, "VEHICULOS")
    por_nuco = collections.defaultdict(list)
    for x in veh.values():
        if num(x.get("NUCCO")):
            por_nuco[num(x["NUCCO"])].append(x)
    fila_de = collections.defaultdict(list)
    for n, r in ver.items():
        fila_de[r.get("ID")].append(n)
    planes, k = [], collections.Counter()
    for d in decisiones(xlsx):
        def no(motivo):
            planes.append(dict(d, fila="", cambios={}, aplica=False, motivo=motivo)); k[motivo.split(":")[0]] += 1
        if not d["nuco"]:
            no("duda: sin NUCO decidido, no se toca"); continue
        if len(fila_de.get(d["id"], [])) != 1:
            no("el ID está %d veces en la hoja" % len(fila_de.get(d["id"], []))); continue
        n = fila_de[d["id"]][0]
        r = ver[n]
        if num(r.get("NUCCO")) != d["nuco_excel"]:
            no("el renglón cambió desde el Excel: NUCCO %s, el Excel decía %s" % (r.get("NUCCO"), d["nuco_excel"])); continue
        vs = por_nuco.get(d["nuco"], [])
        if len(vs) != 1:
            no("el NUCO %s está %d veces en VEHICULOS" % (d["nuco"], len(vs))); continue
        real = vs[0]
        nuevo = {}
        if r.get("ID VEHICULO") != real["ID"]:
            nuevo = {"ID VEHICULO": real["ID"], "FOLIO VEHICULO": real.get("FOLIO", ""), "PLACA": real.get("PLACA", ""), "NUCCO": d["nuco"]}
        elif num(r.get("NUCCO")) != d["nuco"]:
            nuevo = {"NUCCO": d["nuco"]}
        if d["fecha"] and d["fecha"] != (r.get(FECHA) or "").strip()[:10]:
            nuevo[FECHA] = d["fecha"]
        cambios = {c: dict(antes=r.get(c, ""), despues=v) for c, v in nuevo.items() if (r.get(c) or "") != v}
        mismo_dia = [r2["ID"] for n2, r2 in ver.items() if n2 != n and r2.get("ID VEHICULO") == real["ID"]
                     and (r2.get(FECHA) or "")[:10] == (d["fecha"] or (r.get(FECHA) or "")[:10])]
        planes.append(dict(d, fila=n, cambios=cambios, aplica=bool(cambios), real_serie=real.get("SERIE VEHICULO", ""),
                           motivo="" if cambios else "ya estaba así (solo falta mover el comprobante)",
                           aviso=("el NUCO %s ya tiene otra verificación ese día: %s" % (d["nuco"], ", ".join(mismo_dia))) if mismo_dia else ""))
        k["se re-liga" if "ID VEHICULO" in cambios else "solo fecha" if cambios else "ya estaba así"] += 1
        if mismo_dia:
            k["con aviso: otra verificación ese día"] += 1
    D.mkdir(parents=True, exist_ok=True)
    json.dump(dict(xlsx=str(xlsx), letras={c: letra(enc.index(c)) for c in COLS}, plan=planes), open(PLAN, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Re-ligar"
    ws.append(["Fila", "ID", "NUCO hoy", "NUCO que queda", "De dónde", "Serie de ese NUCO", "Fecha nueva", "¿Se escribe?", "Cambios", "Aviso / por qué no"])
    for x in sorted(planes, key=lambda x: (not x["aplica"], x["fuente"], int(x["nuco"] or 0))):
        ws.append([x["fila"], x["id"], x["nuco_excel"], x["nuco"], x["fuente"], x.get("real_serie", ""), x["fecha"], "SÍ" if x["aplica"] else "NO",
                   "; ".join("%s: %r → %r" % (c, v["antes"], v["despues"]) for c, v in x["cambios"].items()),
                   x.get("aviso") or x["motivo"]])
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial", size=10)
            if x.get("aviso"):
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEFGHIJ", [7, 22, 9, 9, 26, 20, 11, 8, 90, 60]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(D / "manual.xlsx")
    print(dict(k))
    print("plan:", PLAN, "| Excel:", D / "manual.xlsx")


def _escribir(h, celdas):
    """celdas: [(columna, rango, valor)]. Las fechas como fecha (USER_ENTERED), lo demás tal cual (RAW)."""
    for opcion, cols in (("RAW", lambda c: c != FECHA), ("USER_ENTERED", lambda c: c == FECHA)):
        data = [{"range": rg, "values": [[v]]} for c, rg, v in celdas if cols(c)]
        for i in range(0, len(data), 300):
            e.llamar(h.values().batchUpdate(spreadsheetId=e.LIBRO, body={"valueInputOption": opcion, "data": data[i:i + 300]}))


def aplicar():
    h = hojas()
    p = json.load(open(PLAN, encoding="utf-8"))
    enc, ver = leer(h, HOJA)
    if any(letra(enc.index(c)) != p["letras"][c] for c in COLS):
        raise SystemExit("Una columna se movió desde el plan")
    bit = e.Bitacora(e.CACHE / "bitacoras" / (time.strftime("%Y%m%d-%H%M%S") + "-religar-manual.jsonl"))
    celdas, k = [], collections.Counter()
    for x in p["plan"]:
        if not x["aplica"]:
            continue
        r = ver.get(x["fila"], {})
        if r.get("ID") != x["id"] or any((r.get(c) or "") != v["antes"] for c, v in x["cambios"].items()):
            k["cambió desde el plan: no se toca"] += 1; continue
        for c, v in x["cambios"].items():
            celdas.append((c, "'%s'!%s%d" % (HOJA, p["letras"][c], x["fila"]), v["despues"]))
            bit.anotar(accion="celda", hoja=HOJA, fila=x["fila"], columna=c, letra=p["letras"][c], id=x["id"], antes=v["antes"],
                       despues=v["despues"], motivo="revisado a mano (%s): NUCO %s (antes %s)" % (x["fuente"], x["nuco"], x["nuco_excel"]))
        k["escrito"] += 1
    _escribir(h, celdas)
    print(dict(k))
    if bit.ruta.exists():
        e.respaldar_bitacora(bit.ruta)
        print("bitácora:", bit.ruta)


def verificar(ruta):
    _, ver = leer(hojas(), HOJA)
    k = collections.Counter()
    for x in (json.loads(l) for l in open(ruta, encoding="utf-8") if l.strip()):
        if x.get("accion") == "celda":
            k["OK" if (ver.get(x["fila"], {}).get(x["columna"]) or "") == x["despues"] else "DISTINTA"] += 1
    print(dict(k))


def deshacer(ruta):
    h = hojas()
    _, ver = leer(h, HOJA)
    celdas, k = [], collections.Counter()
    for x in (json.loads(l) for l in open(ruta, encoding="utf-8") if l.strip()):
        if x.get("accion") != "celda":
            continue
        if (ver.get(x["fila"], {}).get(x["columna"]) or "") != x["despues"]:
            k["cambió después: no se toca"] += 1; continue
        celdas.append((x["columna"], "'%s'!%s%d" % (HOJA, x["letra"], x["fila"]), x["antes"])); k["regresada"] += 1
    _escribir(h, celdas)
    print(dict(k))


if __name__ == "__main__":
    a = sys.argv[1:]
    {"plan": lambda: plan(a[1]), "aplicar": aplicar, "verificar": lambda: verificar(a[1]), "deshacer": lambda: deshacer(a[1])}.get(
        a[0] if a else "", lambda: print(__doc__))()
