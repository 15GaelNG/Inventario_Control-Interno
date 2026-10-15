"""Responsivas de sensores de la app al expediente del NUCO (Fase 3, corrida 3).

    sensores_app.py plan                    SOLO LECTURA → .cache/sensores/plan.json y plan.xlsx
    sensores_app.py aplicar [--nucos 1,2]   EN LA REAL. Bitácora .cache/bitacoras/<fecha>-sensores-app.jsonl (+ Drive)
    sensores_app.py verificar <bitácora>
    sensores_app.py deshacer <bitácora>

Todo lo de ControlVehicular/INSTALACION DE SENSORES_Files_ va a NUCOS VEHICULOS (usuario, 14-oct-2026: "nada ya en donde
estaba antes"), a <NUCO>/1.-DOCUMENTACIÓN/5.-RESPONSIVA/SENSORES, como lo guarda la app (SensoresService):

  - CITADO: el PDF que abre la celda RESPONSIVA SENSOR → SENSORES/RESPONSIVA SENSOR-NNNN <FECHA INSTALACION>.pdf
    ((2) si ya existe), y la celda guarda su liga (mover conserva el ID).
  - VERSIÓN ANTERIOR: AppSheet regeneró el PDF del mismo registro (su prefijo es el ID ANTERIOR del registro) →
    SENSORES/ANTERIORES con su nombre original; si es idéntico (mismo md5) al citado → SENSORES/DUPLICADOS.
  - Su registro no tiene vehículo en VEHICULOS → NUCOS VEHICULOS/_POR REVISAR/SENSORES SIN NUCO, y la celda abre ahí.
  - SIN REGISTRO → _POR REVISAR del NUCO si el PDF trae la placa o la serie del sensor de un solo vehículo; si no,
    NUCOS VEHICULOS/_POR REVISAR/SENSORES SIN NUCO.

El vehículo es el del registro: su ID VEHICULO (y su serie debe ser la del vehículo); sin ID, por la serie. Nunca por
folio. Si el PDF trae la placa o la serie del sensor de OTRO vehículo, no se mueve: REVISAR.
"""
import sys, json, re, time, collections
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e, reglas
import inspecciones_app as ia
from religar_verif import hojas, leer, letra, num

ORIGEN = "1EIMbBdhASdi9RIcO6oPb1FBd72g17-LR"     # ControlVehicular/INSTALACION DE SENSORES_Files_ (prod)
CARPETA_VIEJA = "INSTALACION DE SENSORES_Files_"
HOJA, COL = "INSTALACION DE SENSORES", "RESPONSIVA SENSOR"
RUTA = ["1.-DOCUMENTACIÓN", "5.-RESPONSIVA", "SENSORES"]
PREFIJO = "RESPONSIVA SENSOR"
C = e.CACHE / "sensores"
PLAN = C / "plan.json"
N = reglas.normal
canon = lambda s: re.sub(r"[^A-Z0-9]", "", N(s or ""))


def plan():
    C.mkdir(parents=True, exist_ok=True)
    d = e.drive()
    h = hojas()
    enc, sen = leer(h, HOJA)
    _, veh = leer(h, "VEHICULOS")
    fs = [f for f in ia.listar(d, ORIGEN) if f["mimeType"] != ia.CARP]
    texto_p = C / "texto.json"          # texto de cada PDF (lo bajó la exploración; los que falten se leen como vacíos)
    texto = json.load(open(texto_p, encoding="utf-8")) if texto_p.exists() else {}

    por_id = {x["ID"]: x for x in veh.values() if x.get("ID")}
    por_serie = collections.defaultdict(list)
    placa_nucos, sensor_nucos = collections.defaultdict(set), collections.defaultdict(set)
    for x in veh.values():
        if len(canon(x.get("SERIE VEHICULO"))) == 17:
            por_serie[canon(x["SERIE VEHICULO"])].append(x)
        if num(x.get("NUCCO")) and len(canon(x.get("PLACA"))) >= 6:
            placa_nucos[canon(x["PLACA"])].add(num(x["NUCCO"]))
    for r in sen.values():
        if num(r.get("NUCCO")) and len(canon(r.get("SERIE SENSOR"))) >= 6:
            sensor_nucos[canon(r["SERIE SENSOR"])].add(num(r["NUCCO"]))
    raiz_nuco = {num(x["name"]): x["id"] for x in ia.listar(d, e.NUCOS) if x["mimeType"] == ia.CARP and x["name"].strip().isdigit()}

    def lo_que_dice(f):
        """{NUCO} según la placa o la serie del sensor que trae el texto del PDF (exactas, sin espacios)."""
        t = canon(texto.get(f["id"], ""))
        if not t or t.startswith("ERROR"):
            return set()
        return {n for p, ns in placa_nucos.items() if p in t for n in ns} | {n for s, ns in sensor_nucos.items() if s in t for n in ns}

    def vehiculo(r):
        """(vehículo, motivo de duda). Por ID VEHICULO; sin ID, por la serie (un solo vehículo)."""
        v = por_id.get(r.get("ID VEHICULO", "").strip())
        if v:
            if canon(r.get("SERIE VEHICULO")) and canon(v.get("SERIE VEHICULO")) != canon(r.get("SERIE VEHICULO")):
                return None, "la serie del registro no es la de su ID VEHICULO"
            return v, ""
        vs = por_serie.get(canon(r.get("SERIE VEHICULO")), [])
        if len(vs) == 1:
            return vs[0], ""
        return None, "%d vehículos con la serie %r y sin ID VEHICULO" % (len(vs), r.get("SERIE VEHICULO"))

    cita = collections.defaultdict(list)
    for n, r in sen.items():
        if r.get(COL, "").startswith(CARPETA_VIEJA + "/"):
            cita[r[COL].split("/", 1)[1]].append((n, r))
    por_anterior = {r["ID ANTERIOR"].strip(): (n, r) for n, r in sen.items() if r.get("ID ANTERIOR", "").strip()}
    por_nombre = collections.defaultdict(list)
    for f in fs:
        por_nombre[f["name"]].append(f)

    ops, k = [], collections.Counter()
    tomados = collections.defaultdict(set)      # (nuco, subcarpeta) → nombres que el plan ya puso
    citado_de = {}                              # fila del registro → archivo citado (para comparar md5)

    def nombre_final(nuco, fecha):
        base = "%s-%04d %04d-%02d-%02d" % (PREFIJO, int(nuco), fecha[0], fecha[1], fecha[2])
        nombre, i = base + ".pdf", 2
        while nombre in tomados[(nuco, "")]:
            nombre, i = "%s (%d).pdf" % (base, i), i + 1
        tomados[(nuco, "")].add(nombre)
        return nombre

    def op(f, **x):
        o = dict(archivo=f["id"], de_padre=(f.get("parents") or [""])[0], de_nombre=f["name"], md5=f.get("md5Checksum", ""),
                 modificado=f.get("modifiedTime", ""), confirma=", ".join(sorted(lo_que_dice(f))), **x)
        if not (f.get("capabilities") or {}).get("canMoveItemWithinDrive"):
            o.update(tipo="REVISAR", motivo="la cuenta no puede moverlo", ruta=None)
        ops.append(o)
        k[o["tipo"]] += 1
        return o

    # 1) Lo que cita cada registro
    for n, r in sorted(sen.items()):
        c = r.get(COL, "")
        if not c or c.startswith("http"):
            continue
        nombre = c.split("/", 1)[1] if c.startswith(CARPETA_VIEJA + "/") else None
        archivos = por_nombre.get(nombre, []) if nombre else []
        base = dict(fila=n, registro=r.get("ID", ""), celda_antes=c)
        if len(archivos) != 1:
            ops.append(dict(base, archivo="", de_nombre=nombre or c, tipo="REVISAR", ruta=None, motivo="%d archivos con ese nombre" % len(archivos)))
            k["REVISAR"] += 1; continue
        f = archivos[0]
        v, duda = vehiculo(r)
        fecha = ia.fecha_de(r.get("FECHA INSTALACION"))
        nuco = num(v.get("NUCCO")) if v else ""
        dice = lo_que_dice(f)
        if not v:
            # Nada se queda en la carpeta vieja: a _POR REVISAR de la raíz, y la celda abre ahí
            op(f, **base, tipo="SIN NUCO", nuco="", ruta=["_POR REVISAR", "SENSORES SIN NUCO"], a_nombre=f["name"],
               celda_nueva=ia.liga(f["id"]), motivo=duda); continue
        if num(r.get("NUCCO")) and num(r.get("NUCCO")) != nuco:
            op(f, **base, tipo="REVISAR", ruta=None, nuco=nuco, motivo="el registro dice NUCCO %s y su vehículo es %s" % (r.get("NUCCO"), nuco)); continue
        if nuco not in raiz_nuco:
            op(f, **base, tipo="REVISAR", ruta=None, nuco=nuco, motivo="NUCO %r sin carpeta en NUCOS VEHICULOS" % nuco); continue
        if dice and nuco not in dice:
            op(f, **base, tipo="REVISAR", ruta=None, nuco=nuco, motivo="el PDF trae la placa o el sensor de NUCO %s" % ", ".join(sorted(dice))); continue
        if not fecha:
            op(f, **base, tipo="REVISAR", ruta=None, nuco=nuco, motivo="sin FECHA INSTALACION"); continue
        o = op(f, **base, tipo="CITADO", nuco=nuco, ruta=RUTA, a_nombre=nombre_final(nuco, fecha), celda_nueva=ia.liga(f["id"]),
               liga_vehiculo="" if r.get("ID VEHICULO") else "sin ID VEHICULO: por la serie",
               motivo="lo cita %s (instalado %s)" % (r.get("ID"), r.get("FECHA INSTALACION")))
        if o["tipo"] == "CITADO":
            citado_de[n] = o

    # 2) Lo que ningún registro cita
    hechos = {o["archivo"] for o in ops}
    for f in sorted(fs, key=lambda f: f["name"]):
        if f["id"] in hechos:
            continue
        pre = f["name"].split(".")[0]
        if pre in por_anterior:
            n, r = por_anterior[pre]
            bueno = citado_de.get(n)
            if bueno:
                sub = "DUPLICADOS" if f.get("md5Checksum") and f.get("md5Checksum") == bueno["md5"] else "ANTERIORES"
                op(f, tipo="VERSIÓN ANTERIOR", nuco=bueno["nuco"], registro=r["ID"], ruta=RUTA + [sub], a_nombre=f["name"],
                   motivo="versión anterior de la responsiva de %s%s" % (r["ID"], " (idéntica)" if sub == "DUPLICADOS" else ""))
                continue
            v, duda = vehiculo(r)
            nuco = num(v.get("NUCCO")) if v else ""
            if v and nuco in raiz_nuco and not r.get(COL):
                op(f, tipo="SIN CITA", nuco=nuco, registro=r["ID"], ruta=RUTA + ["_POR REVISAR"], a_nombre=f["name"],
                   motivo="es de %s, pero su celda está vacía" % r["ID"])
                continue
            op(f, tipo="REVISAR", nuco=nuco, registro=r["ID"], ruta=None, motivo="es de %s, cuya responsiva no se mueve: %s" % (r["ID"], duda or "revisar"))
            continue
        dice = lo_que_dice(f)
        if len(dice) == 1 and next(iter(dice)) in raiz_nuco:
            nuco = next(iter(dice))
            op(f, tipo="SIN REGISTRO", nuco=nuco, ruta=RUTA + ["_POR REVISAR"], a_nombre=f["name"],
               motivo="ningún registro lo cita; la placa o el sensor dicen NUCO %s" % nuco)
        else:
            op(f, tipo="SIN REGISTRO", nuco="", ruta=["_POR REVISAR", "SENSORES SIN NUCO"], a_nombre=f["name"],
               motivo="ningún registro lo cita y no se lee el vehículo" + (" (dice %s)" % sorted(dice) if dice else ""))

    # Nombres repetidos dentro de una misma subcarpeta del plan → (2)
    vistos = collections.defaultdict(set)
    for o in ops:
        if o.get("ruta") and o["tipo"] != "CITADO":
            clave_c = (o.get("nuco"), tuple(o["ruta"]))
            nombre, i = o["a_nombre"], 2
            while nombre in vistos[clave_c]:
                nombre, i = "%s (%d)%s" % (Path(o["a_nombre"]).stem, i, Path(o["a_nombre"]).suffix), i + 1
            o["a_nombre"] = nombre
            vistos[clave_c].add(nombre)

    json.dump(dict(letra=letra(enc.index(COL)), ops=ops), open(PLAN, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Plan"
    ws.append(["Qué es", "Registro", "Fila", "NUCO", "Archivo de hoy", "Va a", "Nombre nuevo", "La celda queda", "El PDF dice NUCO", "Por qué", "Abrir"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    color = {"SIN NUCO": "FFF2CC", "CITADO": "E2EFDA", "VERSIÓN ANTERIOR": "F2F2F2", "SIN REGISTRO": "FFF2CC", "SIN CITA": "FFF2CC", "REVISAR": "F8CBAD"}
    for o in sorted(ops, key=lambda o: (["REVISAR", "SIN NUCO", "SIN REGISTRO", "SIN CITA", "CITADO", "VERSIÓN ANTERIOR"].index(o["tipo"]), str(o.get("nuco")))):
        ruta = o.get("ruta")
        ws.append([o["tipo"], o.get("registro", ""), o.get("fila", ""), o.get("nuco", ""), o["de_nombre"],
                   ("NUCOS VEHICULOS/" + "/".join(ruta) if ruta and ruta[0] == "_POR REVISAR" else "%s/%s" % (o.get("nuco"), "/".join(ruta))) if ruta else "(no se mueve)",
                   o.get("a_nombre", ""), "liga al archivo" if o.get("celda_nueva") else "", o.get("confirma", ""), o["motivo"],
                   "https://drive.google.com/file/d/%s/view" % o["archivo"] if o["archivo"] else ""])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=color[o["tipo"]])
    for col, w in zip("ABCDEFGHIJK", [16, 20, 6, 6, 44, 60, 40, 14, 12, 60, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(C / "plan.xlsx")
    citados = [o for o in ops if o["tipo"] == "CITADO"]
    print(len(fs), "archivos en", CARPETA_VIEJA, "|", dict(k))
    print("citados: el PDF confirma su NUCO en", sum(1 for o in citados if o["confirma"]), "| no trae placa ni sensor legible en",
          sum(1 for o in citados if not o["confirma"]))
    print("anteriores:", collections.Counter(o["ruta"][-1] for o in ops if o["tipo"] == "VERSIÓN ANTERIOR"))
    print("NUCO distintos:", len({o.get("nuco") for o in ops if o.get("ruta") and o.get("nuco")}))
    for o in ops:
        if o["tipo"] in ("REVISAR", "SIN NUCO", "SIN REGISTRO", "SIN CITA"):
            print("  ", o["tipo"], o.get("registro", ""), o.get("nuco", ""), o["de_nombre"], "→", o["motivo"])
    print("plan:", PLAN, "| excel:", C / "plan.xlsx")


def aplicar(nucos=None):
    d = e.drive()
    h = hojas()
    p = json.load(open(PLAN, encoding="utf-8"))
    bit = e.Bitacora(e.CACHE / "bitacoras" / (time.strftime("%Y%m%d-%H%M%S") + "-sensores-app.jsonl"))
    k = collections.Counter()
    carpetas, listados = {}, {}

    def carpeta(padre, nombre, nuco):
        """La subcarpeta por su nombre (como ExpedienteNuco.subcarpeta_: exacta, o igual sin espacios ni acentos); si no
        existe, se crea. Si hay dos, no se escoge."""
        if (padre, nombre) in carpetas:
            return carpetas[(padre, nombre)]
        dentro = [x for x in ia.listar(d, padre) if x["mimeType"] == ia.CARP]
        ya = [x for x in dentro if x["name"] == nombre] or [x for x in dentro if ia.clave(x["name"]) == ia.clave(nombre)]
        if len(ya) > 1:
            raise RuntimeError("hay %d carpetas %r en %s" % (len(ya), nombre, padre))
        if ya:
            i = ya[0]["id"]
        else:
            i = e.llamar(d.files().create(body={"name": nombre, "mimeType": ia.CARP, "parents": [padre]}, fields="id", supportsAllDrives=True))["id"]
            bit.anotar(accion="crear_carpeta", nuco=nuco, id=i, padre=padre, nombre=nombre)
            k["carpetas creadas"] += 1
        carpetas[(padre, nombre)] = i
        return i

    raiz_nuco = {num(x["name"]): x["id"] for x in ia.listar(d, e.NUCOS) if x["mimeType"] == ia.CARP and x["name"].strip().isdigit()}
    _, sen = leer(h, HOJA)
    for o in p["ops"]:
        if not o.get("ruta") or (nucos and str(o.get("nuco")) not in nucos):
            continue
        f = e.llamar(d.files().get(fileId=o["archivo"], fields="name,parents,trashed", supportsAllDrives=True))
        if f.get("trashed") or (f.get("parents") or [""])[0] != o["de_padre"] or f["name"] != o["de_nombre"]:
            bit.anotar(accion="saltado", nuco=o.get("nuco", ""), id=o["archivo"], motivo="ya no está donde lo vio el plan"); k["saltados"] += 1; continue
        ruta = o["ruta"]
        destino = e.NUCOS if ruta[0] == "_POR REVISAR" else raiz_nuco[o["nuco"]]
        for nombre in ruta:
            destino = carpeta(destino, nombre, o.get("nuco", ""))
        if destino not in listados:
            listados[destino] = {x["name"] for x in ia.listar(d, destino)}
        nombre, i = o["a_nombre"], 2
        while nombre in listados[destino]:
            nombre, i = "%s (%d)%s" % (Path(o["a_nombre"]).stem, i, Path(o["a_nombre"]).suffix), i + 1
        listados[destino].add(nombre)
        e.llamar(d.files().update(fileId=o["archivo"], addParents=destino, removeParents=o["de_padre"], fields="id", supportsAllDrives=True,
                                  body={"name": nombre, "appProperties": {"origen_padre": o["de_padre"], "origen_nombre": o["de_nombre"][:100]}}))
        bit.anotar(accion="mover", nuco=o.get("nuco", ""), id=o["archivo"], de_padre=o["de_padre"], de_nombre=o["de_nombre"], a_padre=destino,
                   a_nombre=nombre, app=True)
        k["movidos"] += 1
        if o.get("celda_nueva"):
            r = sen.get(o["fila"], {})
            if r.get("ID") == o["registro"] and r.get(COL) == o["celda_antes"]:
                e.llamar(h.values().update(spreadsheetId=e.LIBRO, range="'%s'!%s%d" % (HOJA, p["letra"], o["fila"]), valueInputOption="RAW",
                                           body={"values": [[o["celda_nueva"]]]}))
                bit.anotar(accion="celda", nuco=o["nuco"], hoja=HOJA, fila=o["fila"], letra=p["letra"], id_registro=o["registro"],
                           antes=o["celda_antes"], despues=o["celda_nueva"])
                k["celdas"] += 1
            else:
                bit.anotar(accion="saltado", nuco=o["nuco"], id=o["archivo"], motivo="la celda ya no dice lo que vio el plan (archivo sí movido)")
                k["celdas saltadas"] += 1
    print(dict(k))
    if bit.ruta.exists():
        e.respaldar_bitacora(bit.ruta)
        print("bitácora:", bit.ruta)


def verificar(ruta):
    d = e.drive()
    _, sen = leer(hojas(), HOJA)
    k = collections.Counter()
    for x in (json.loads(l) for l in open(ruta, encoding="utf-8") if l.strip()):
        if x["accion"] == "mover":
            f = e.llamar(d.files().get(fileId=x["id"], fields="name,parents,trashed", supportsAllDrives=True))
            k["archivo en su lugar" if not f.get("trashed") and (f.get("parents") or [""])[0] == x["a_padre"] and f["name"] == x["a_nombre"] else "DISTINTO"] += 1
        elif x["accion"] == "celda":
            k["celda OK" if sen.get(x["fila"], {}).get(COL) == x["despues"] else "CELDA DISTINTA"] += 1
        else:
            k[x["accion"]] += 1
    print(dict(k))
    print("quedan en", CARPETA_VIEJA + ":", sum(1 for f in ia.listar(d, ORIGEN) if f["mimeType"] != ia.CARP))


def deshacer(ruta):
    h = hojas()
    _, sen = leer(h, HOJA)
    for x in (json.loads(l) for l in open(ruta, encoding="utf-8") if l.strip()):
        if x["accion"] == "celda" and sen.get(x["fila"], {}).get(COL) == x["despues"]:
            e.llamar(h.values().update(spreadsheetId=e.LIBRO, range="'%s'!%s%d" % (HOJA, x["letra"], x["fila"]), valueInputOption="RAW",
                                       body={"values": [[x["antes"]]]}))
    e.deshacer(ruta)


if __name__ == "__main__":
    a = sys.argv[1:]
    nucos = a[a.index("--nucos") + 1].split(",") if "--nucos" in a else None
    {"plan": plan, "aplicar": lambda: aplicar(nucos), "verificar": lambda: verificar(a[1]), "deshacer": lambda: deshacer(a[1])}.get(
        a[0] if a else "", lambda: print(__doc__))()
