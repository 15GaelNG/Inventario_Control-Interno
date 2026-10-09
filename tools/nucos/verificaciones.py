"""Homologar `3.- VERIFICACIONES` de cada NUCO (docs/nucos-verificaciones.md). SOLO LECTURA: no mueve nada.

  foto.py <salida> VERIFIC          la foto de la carpeta (carpetas y archivos con md5) → .cache/verificaciones/foto.json
  verificaciones.py ocr             lee cada comprobante con el OCR de Google: una copia como Documento en una carpeta
                                    temporal propia (fuera de NUCOS), se baja el texto y se borra la copia. Se guarda en
                                    .cache/verificaciones/ocr.json y se puede cortar y seguir.
  verificaciones.py dry-run <xlsx>  el plan de cada archivo (→ plan.json) y un Excel para revisarlo.

El periodo sale de, en este orden: la fecha de expedición del certificado (OCR), la fecha del nombre del archivo
(`B-153 VERIFICACION 2026 04 27.pdf`), la carpeta del área (`PROGRAMA 2025/PRIMER SEMESTRE`; sus "BIMESTRE" también son
semestres, lo dicen los certificados), el registro de la app si el NUCO tiene uno solo en un semestre. Nunca la fecha
de subida. Lo que no se puede fechar va a `_POR REVISAR`.
"""
import sys, json, re, collections, datetime, threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import expediente as e, reglas

CACHE = Path(__file__).resolve().parent / ".cache" / "verificaciones"
FOTO, OCR, APP, PLAN = CACHE / "foto.json", CACHE / "ocr.json", CACHE / "app.json", CACHE / "plan.json"
FOTO_DOC = Path(__file__).resolve().parent / ".cache" / "inspecciones" / "foto-doc.json"
RAIZ = "3.- VERIFICACIONES"
POR_REVISAR = "_POR REVISAR"
SEMESTRE = {1: "1ER SEMESTRE", 2: "2DO SEMESTRE"}
LEIBLES = re.compile(r"^(image/|application/pdf)")
TEMPORAL = "VERIFICACIONES OCR TEMPORAL"
AJENO = re.compile(r"ANTIGUA|SINIESTRO|CONVENIO|EXPEDIENTE|CORRALON|ACCIDENTE|TALLER|SERVICIO")
VERSION_ANTERIOR = {"RESPONSIVA": "RESPONSIVAS ANTERIORES", "SEGURO": "SEGUROS ANTERIORES", "TENENCIA": "TENENCIAS ANTERIORES"}
N = reglas.normal


# ---------------------------------------------------------------- OCR

def ocr():
    foto = json.load(open(FOTO, encoding="utf-8"))
    hechos = json.load(open(OCR, encoding="utf-8")) if OCR.exists() else {}
    faltan = [a for a in foto["archivos"] if LEIBLES.match(a["mime"]) and a["id"] not in hechos]
    print("a leer: %d (ya leídos %d)" % (len(faltan), len(hechos)))
    d = e.drive()
    tmp = e.llamar(d.files().create(body={"name": TEMPORAL, "mimeType": reglas.CARPETA}, fields="id"))["id"]
    candado = threading.Lock()
    local = threading.local()

    def uno(a):
        if not hasattr(local, "d"):
            local.d = e.drive()
        dd = local.d
        try:
            c = e.llamar(dd.files().copy(fileId=a["id"], ocrLanguage="es", supportsAllDrives=True, fields="id",
                                         body={"name": "OCR " + a["id"], "mimeType": "application/vnd.google-apps.document",
                                               "parents": [tmp]}))
            try:
                texto = e.llamar(dd.files().export(fileId=c["id"], mimeType="text/plain")).decode("utf-8", "replace")
            finally:
                e.llamar(dd.files().delete(fileId=c["id"]))
            r = dict(texto=texto.lstrip("﻿"))
        except Exception as ex:
            r = dict(error=str(ex)[:300])
        with candado:
            hechos[a["id"]] = r
            if len(hechos) % 20 == 0:
                json.dump(hechos, open(OCR, "w", encoding="utf-8"), ensure_ascii=False)
                print("  %d/%d" % (len(hechos), len(hechos) + len(faltan) - sum(1 for x in faltan if x["id"] in hechos)))

    def de_la_imagen(a):
        """PDF escaneado con una capa de texto en blanco: Google no lo pasa por OCR (regresa vacío). Se saca la imagen
        de la página 1 y esa se sube (a la carpeta temporal) como Documento con OCR."""
        import io
        from pypdf import PdfReader
        from googleapiclient.http import MediaIoBaseUpload
        if not hasattr(local, "d"):
            local.d = e.drive()
        dd = local.d
        try:
            pagina = PdfReader(io.BytesIO(e.llamar(dd.files().get_media(fileId=a["id"], supportsAllDrives=True)))).pages[0]
            if not pagina.images:
                return
            img = pagina.images[0]
            mime = "image/png" if img.name.lower().endswith(".png") else "image/jpeg"
            c = e.llamar(dd.files().create(body={"name": "OCR " + a["id"], "mimeType": "application/vnd.google-apps.document",
                                                 "parents": [tmp]}, ocrLanguage="es", fields="id",
                                           media_body=MediaIoBaseUpload(io.BytesIO(img.data), mimetype=mime)))
            try:
                texto = e.llamar(dd.files().export(fileId=c["id"], mimeType="text/plain")).decode("utf-8", "replace")
            finally:
                e.llamar(dd.files().delete(fileId=c["id"]))
            with candado:
                hechos[a["id"]] = dict(texto=texto.lstrip("﻿"), de="imagen del PDF")
        except Exception as ex:
            with candado:
                hechos[a["id"]] = dict(texto="", error="imagen del PDF: " + str(ex)[:250])

    try:
        with ThreadPoolExecutor(4) as ex:
            list(ex.map(uno, faltan))
        vacios = [a for a in foto["archivos"] if a["mime"] == "application/pdf" and not hechos.get(a["id"], {}).get("texto", "").strip()
                  and hechos.get(a["id"], {}).get("de") != "imagen del PDF"]
        print("PDF con texto vacío, por su imagen: %d" % len(vacios))
        with ThreadPoolExecutor(4) as ex:
            list(ex.map(de_la_imagen, vacios))
    finally:
        json.dump(hechos, open(OCR, "w", encoding="utf-8"), ensure_ascii=False)
        e.llamar(d.files().delete(fileId=tmp))
    print("leídos %d | con error %d" % (len(hechos), sum(1 for x in hechos.values() if "error" in x)))


# ---------------------------------------------------------------- lectura del certificado

def lee_certificado(texto):
    """{fecha, motivo, placas, certificado} de lo que el OCR alcanzó a leer."""
    t = texto or ""
    hoy = datetime.date.today()
    fechas = []
    for d_, m_, y_ in re.findall(r"(?<!\d)(\d{1,2})/(\d{1,2})/(20\d\d)(?!\d)", t):
        try:
            f = datetime.date(int(y_), int(m_), int(d_))
        except ValueError:
            continue
        if datetime.date(2015, 1, 1) <= f <= hoy:
            fechas.append(f)
    T = N(t)
    motivo = next((m for m, rx in (("RECHAZO", r"RECHAZ"), ("SANCION", r"SANCI"), ("NORMAL", r"\bNORMA"))
                   if re.search(rx, T)), "")
    certificado = bool(re.search(r"VERIFICACI|CERTIFICADO|SEMESTRE", T))
    placas = set(re.findall(r"\b[A-Z]{2,3}\d{3,4}[A-Z]\b|\b[A-Z]\d{2}[A-Z]{3}\b", T))
    fecha, corregida = (fechas[0] if fechas else None), False
    # El folio del certificado empieza con el año (2620138468 → 2026): corrige un año mal leído (2020 por 2026)
    # Solo si todas las lecturas del folio dicen el mismo año y la vigencia ("Hasta Enero/Febrero del 2027", el mismo
    # año o el siguiente) no le da la razón a la fecha leída
    folios = {int(x) for x in re.findall(r"(?<!\d)(2\d)\d{8}(?!\d)", t)}
    vig = re.search(r"VIGENCIA[^\n]*?(20\d\d)", T)
    cuadra = lambda anio: not vig or 0 <= int(vig.group(1)) - anio <= 1
    if fecha and certificado and len(folios) == 1 and (not cuadra(fecha.year) if vig else 2000 + min(folios) != fecha.year):
        anio = 2000 + folios.pop()
        if anio != fecha.year and 2015 <= anio <= hoy.year and cuadra(anio):
            try:
                fecha, corregida = fecha.replace(year=anio), True
            except ValueError:
                pass
    return dict(fecha=fecha.isoformat() if fecha else "", motivo=motivo, certificado=certificado, placas=placas,
                corregida=corregida, doble_cero=bool(re.search(r"DOBLE CERO", T)))


def fecha_del_nombre(nombre):
    m = re.search(r"(?<!\d)(20\d\d)[ _.-]?(\d\d)[ _.-]?(\d\d)(?!\d)", nombre)
    if m:
        try:
            return datetime.date(*map(int, m.groups())).isoformat()
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d\d)[ _.-](\d\d)[ _.-](20\d\d)(?!\d)", nombre)
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            pass
    return ""


def periodo_de_carpeta(ruta):
    """(año, semestre) de la carpeta del área: PROGRAMA 2025/PRIMER SEMESTRE, VERIFICACION 2026/1 SEMESTRE…
    PRIMER/SEGUNDO BIMESTRE son el 1er/2do semestre (los certificados de esas carpetas lo dicen)."""
    anio, sem = "", 0
    for s in ruta:
        n = N(s)
        m = re.search(r"(?<!\d)(20\d\d)(?!\d)", n)
        if m:
            anio = m.group(1)
        if re.search(r"SEMESTRE|BIMESTRE|SEMETRE", n):
            if re.search(r"PRIMER|\b1\b|1ER|1RO", n):
                sem = 1
            elif re.search(r"SEGUND|\b2\b|2DO", n):
                sem = 2
    return anio, sem


# ---------------------------------------------------------------- plan

def dry_run(salida):
    foto = json.load(open(FOTO, encoding="utf-8"))
    leido = json.load(open(OCR, encoding="utf-8")) if OCR.exists() else {}
    app = json.load(open(APP, encoding="utf-8"))
    doc = json.load(open(FOTO_DOC, encoding="utf-8")) if FOTO_DOC.exists() else {"archivos": []}
    nuco_num = lambda s: str(int(re.match(r"\s*(\d+)", s).group(1))) if re.match(r"\s*(\d+)", s) else s

    # El registro de la app por NUCO: (año, semestre) de cada verificación
    app_sem = collections.defaultdict(set)
    for v in app["ver"]:
        m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", v.get("FECHA VERIFICACION") or "")
        nuco = app["nuco_de"].get(v.get("FOLIO VEHICULO"), "")
        if m and nuco:
            app_sem[nuco_num(nuco)].add((m.group(3), 1 if int(m.group(2)) <= 6 else 2))
    placa_de = collections.defaultdict(set)
    for v in app["ver"]:
        nuco = app["nuco_de"].get(v.get("FOLIO VEHICULO"), "")
        if nuco and v.get("PLACA"):
            placa_de[nuco_num(nuco)].add(N(v["PLACA"]).replace(" ", ""))
    md5_doc = collections.defaultdict(set)
    for d in doc["archivos"]:
        md5_doc[nuco_num(d["nuco"])].add(d.get("md5"))

    plan = []
    for a in foto["archivos"]:
        nuco = nuco_num(a["nuco"])
        ruta = a["ruta"][1:-1]                          # entre la raíz y el archivo
        texto_ruta = N(" / ".join(ruta))
        r = leido.get(a["id"], {})
        c = lee_certificado(r.get("texto", ""))
        base = dict(id=a["id"], nuco=a["nuco"], antes="/".join(a["ruta"]), mime=a["mime"], md5=a.get("md5"),
                    ocr_fecha=c["fecha"], ocr_motivo=c["motivo"], ocr_error=r.get("error", ""),
                    es_certificado="SI" if c["certificado"] else "", placa_ok="", fuente="", accion="", despues="", aviso="")
        if placa_de.get(nuco) and c["placas"]:
            base["placa_ok"] = "SI" if placa_de[nuco] & c["placas"] else "NO: " + ", ".join(sorted(c["placas"]))[:60]

        if a["mime"] == "application/vnd.google-apps.shortcut":
            plan.append(dict(base, accion="REVISAR", aviso="acceso directo"))
            continue
        # Lo que no es una verificación
        if AJENO.search(texto_ruta):
            plan.append(dict(base, accion="REVISAR", aviso="no es verificación (expediente viejo, siniestro, taller…)"))
            continue
        conc = reglas.concepto_archivo(a["name"]) or next(
            (reglas.concepto(s) for s in reversed(ruta) if reglas.concepto(s)), None)
        if conc and not re.search(r"VERIFIC", N(a["name"])):
            if a.get("md5") and a["md5"] in md5_doc[nuco]:
                plan.append(dict(base, accion="PAPELERA", aviso="ya está igual (mismo md5) en 1.-DOCUMENTACIÓN"))
                continue
            carpeta = reglas.CARPETA_DE[conc]
            rel = "/".join([carpeta] + ([VERSION_ANTERIOR[conc]] if conc in VERSION_ANTERIOR else []))
            m = re.search(r"(?<!\d)(20[0-2]\d)(?!\d)", a["name"] + " " + " ".join(ruta))
            fecha = (m.group(1) if m else "")
            plan.append(dict(base, accion="MOVER", destino_tipo="DOCUMENTACION", fuente="nombre",
                             despues="%s/%s/%s/@%s|%s|%s" % (a["nuco"], reglas.DOCUMENTACION, rel, reglas.PREFIJO_DE[conc], fecha,
                                                            reglas.extension(a["name"], a["mime"])),
                             aviso="documento %s guardado en verificaciones" % conc))
            continue

        # El periodo
        fecha, fuente = c["fecha"], "OCR del certificado" if c["fecha"] else ""
        if c["corregida"]:
            base["aviso"] = "año del OCR corregido con el folio del certificado"
        if not fecha:
            fecha = fecha_del_nombre(a["name"])
            fuente = "fecha del nombre" if fecha else ""
        anio, sem = (fecha[:4], 1 if int(fecha[5:7]) <= 6 else 2) if fecha else ("", 0)
        if not fecha:
            anio, sem = periodo_de_carpeta(ruta)
            if anio and sem:
                fuente = "carpeta del área"
            elif len(app_sem.get(nuco, ())) == 1:
                anio, sem = next(iter(app_sem[nuco]))
                fuente = "único registro de la app"
            else:
                anio, sem = "", 0
        if not (anio and sem):
            plan.append(dict(base, accion="REVISAR", aviso="no se pudo saber el periodo (sin fecha legible ni carpeta)"))
            continue
        doble = c["doble_cero"] or "DOBLE CERO" in texto_ruta or "DOBLE CERO" in N(a["name"])
        sufijo = (" DOBLE CERO" if doble else "") + (" RECHAZO" if c["motivo"] == "RECHAZO" else "")
        plan.append(dict(base, accion="MOVER", destino_tipo="PERIODO", fuente=fuente,
                         despues="%s/%s/%s/%s/@VERIFICACION|%s|%s|%s" % (a["nuco"], RAIZ, anio, SEMESTRE[sem], fecha, sufijo,
                                                                         reglas.extension(a["name"], a["mime"]))))

    # _POR REVISAR: plano, el nombre lleva de dónde venía
    for p in plan:
        if p["accion"] == "REVISAR":
            a = next(x for x in foto["archivos"] if x["id"] == p["id"])
            origen = next((s.strip() for s in reversed(a["ruta"][1:-1])), "")
            p.update(accion="MOVER", destino_tipo="POR REVISAR",
                     despues="/".join([a["nuco"], RAIZ, POR_REVISAR, (origen + " - " if origen else "") + a["name"]]))

    # Nombres finales: VERIFICACION-0088 2026-06-11 RECHAZO.pdf; idéntico (mismo md5) en el mismo destino → papelera
    usados = collections.defaultdict(set)
    vistos = collections.defaultdict(set)
    for p in sorted(plan, key=lambda p: (p["despues"], p["antes"])):
        if p["accion"] != "MOVER":
            continue
        carpeta, _, hoja = p["despues"].rpartition("/")
        if p.get("md5") and p["md5"] in vistos[carpeta]:
            p.update(accion="PAPELERA", despues="", aviso="copia idéntica (mismo md5) de otro archivo que va a la misma carpeta")
            continue
        vistos[carpeta].add(p.get("md5"))
        if not hoja.startswith("@"):
            nombre = hoja
            n = 2
            while nombre in usados[carpeta]:
                b, x = re.match(r"(.*?)(\.[^.]+)?$", hoja).groups()
                nombre = "%s (%d)%s" % (b, n, x or "")
                n += 1
        else:
            partes = hoja[1:].split("|")
            prefijo, fecha, ext = partes[0], partes[1], partes[-1]
            sufijo = partes[2] if len(partes) == 4 else ""
            b = "%s-%04d" % (prefijo, int(nuco_num(p["nuco"]))) + (" " + fecha if fecha else "") + sufijo
            nombre, n = b + ext, 2
            while nombre in usados[carpeta]:
                nombre = "%s (%d)%s" % (b, n, ext)
                n += 1
        usados[carpeta].add(nombre)
        p["despues"] = carpeta + "/" + nombre

    # Las raíces con otro nombre y los NUCO sin carpeta
    renombrar = [r for r in foto["raices"] if r["name"] != RAIZ]
    con = {r["nuco"] for r in foto["raices"]}
    sin = sorted((n["name"] for n in foto["nucos"] if n["name"] not in con), key=lambda s: int(nuco_num(s)) if nuco_num(s).isdigit() else 9999)
    json.dump(dict(plan=plan, renombrar=renombrar, sin_carpeta=sin), open(PLAN, "w", encoding="utf-8"), ensure_ascii=False)
    excel(salida, plan, renombrar, sin, foto)


def excel(salida, plan, renombrar, sin, foto):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    neg = Font(name="Arial", bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F3A5F")
    normal = Font(name="Arial", size=10)

    def hoja(ws, enc, filas, anchos):
        ws.append(enc)
        for c in ws[1]:
            c.font, c.fill = neg, fondo
        for f in filas:
            ws.append(f)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.font = normal
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Resumen"
    cuenta = collections.Counter((p["accion"], p.get("destino_tipo", "")) for p in plan)
    fuentes = collections.Counter(p["fuente"] for p in plan if p.get("destino_tipo") == "PERIODO")
    filas = [["Archivos en 3.- VERIFICACIONES", len(plan)],
             ["NUCO con carpeta / sin carpeta", "%d / %d" % (len(foto["raices"]), len(sin))],
             ["Carpetas raíz a renombrar a '%s'" % RAIZ, len(renombrar)], ["", ""]]
    filas += [["%s → %s" % (a, t or "-"), n] for (a, t), n in sorted(cuenta.items())]
    filas += [["", ""], ["El periodo salió de…", ""]] + [["   " + f, n] for f, n in fuentes.most_common()]
    filas += [["", ""], ["Certificados con RECHAZO", sum(1 for p in plan if "RECHAZO" in p["despues"])],
              ["Certificados cuya placa NO es la del vehículo (según la app)", sum(1 for p in plan if p["placa_ok"].startswith("NO"))]]
    hoja(ws, ["Concepto", "Cantidad"], filas, [62, 14])

    orden = {"POR REVISAR": 0, "DOCUMENTACION": 1, "PERIODO": 2, "": 3}
    ws = wb.create_sheet("Archivos")
    filas = [[p["nuco"], p["accion"], p.get("destino_tipo", ""), p["antes"], p["despues"], p["fuente"], p["ocr_fecha"],
              p["ocr_motivo"], p["es_certificado"], p["placa_ok"], p["aviso"] or p["ocr_error"],
              '=HYPERLINK("https://drive.google.com/file/d/%s/view","abrir")' % p["id"]]
             for p in sorted(plan, key=lambda p: (orden.get(p.get("destino_tipo", ""), 9),
                                                  int(re.match(r"\d+", p["nuco"]).group()) if re.match(r"\d+", p["nuco"]) else 0))]
    hoja(ws, ["NUCO", "Acción", "Destino", "Ruta de antes", "Ruta nueva", "Periodo por", "Fecha OCR", "Motivo OCR",
              "¿Certificado?", "¿Placa del vehículo?", "Aviso", "Archivo"], filas, [8, 10, 14, 60, 60, 22, 12, 11, 12, 22, 45, 9])

    ws = wb.create_sheet("Renombrar raíz")
    hoja(ws, ["NUCO", "Nombre de antes", "Nombre nuevo"], [[r["nuco"], r["name"], RAIZ] for r in renombrar], [8, 34, 24])
    ws = wb.create_sheet("NUCO sin carpeta")
    hoja(ws, ["NUCO (no tiene 3.- VERIFICACIONES)"], [[s] for s in sin], [40])
    wb.save(salida)
    print("Excel:", salida)
    for (a, t), n in sorted(cuenta.items()):
        print("  %-10s %-14s %d" % (a, t or "-", n))
    print("  periodo por:", dict(fuentes))


if __name__ == "__main__":
    if sys.argv[1:2] == ["ocr"]:
        ocr()
    elif sys.argv[1:2] == ["dry-run"]:
        dry_run(sys.argv[2])
    else:
        print(__doc__)
